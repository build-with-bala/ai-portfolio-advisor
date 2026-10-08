"""Export the HTML deck (landing/deck/) to docs/Capstone_Deck.pdf and docs/Capstone_Deck.pptx.

    .venv/bin/python scripts/export_deck.py                 # capture with headless Chrome, then build both files
    .venv/bin/python scripts/export_deck.py --skip-capture  # rebuild the PPTX from PNGs already in landing/deck/export/

How it works
  * A throwaway HTTP server serves landing/ on a free local port.
  * Headless Chrome opens deck/?export=N once per slide (final, fully revealed state, 1920x1080)
    and saves landing/deck/export/slide-NN.png.
  * Headless Chrome prints deck/?print (one slide per 1920x1080 px page) to a vector PDF.
  * python-pptx places each PNG full-bleed on a 16:9 slide and writes the slide title and key
    text into the speaker notes, so the file stays searchable.

Needs: python-pptx and Pillow, a Chromium browser (Playwright's chrome-headless-shell if present,
otherwise Google Chrome or Chromium; set CHROME_BIN to choose one), and network access for the
two Google Fonts.
"""
from __future__ import annotations

import argparse
import functools
import glob
import http.server
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
LANDING = ROOT / "landing"
DECK = LANDING / "deck"
PNG_DIR = DECK / "export"
PDF_OUT = ROOT / "docs" / "Capstone_Deck.pdf"
PPTX_OUT = ROOT / "docs" / "Capstone_Deck.pptx"
LIVE_URL = "https://fin.iimbg.com/deck/"
W, H = 1920, 1080
EMU_PER_PX = 9525  # 96 dpi: 1920 x 1080 px = 20 x 11.25 in, a 16:9 slide


class SlideMeta(HTMLParser):
    """Collects data-title / data-notes from each <section class="slide">."""

    def __init__(self) -> None:
        super().__init__()
        self.slides: list[dict[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        a = dict(attrs)
        if tag == "section" and "slide" in (a.get("class") or "").split():
            self.slides.append({
                "title": a.get("data-title") or "",
                "section": a.get("data-section") or "",
                "notes": a.get("data-notes") or "",
            })


def read_slides() -> list[dict[str, str]]:
    parser = SlideMeta()
    parser.feed((DECK / "index.html").read_text(encoding="utf-8"))
    if not parser.slides:
        sys.exit("No <section class=\"slide\"> found in landing/deck/index.html")
    return parser.slides


def find_chrome() -> str:
    shells = sorted(glob.glob(os.path.expanduser(
        "~/Library/Caches/ms-playwright/chromium_headless_shell-*/chrome-headless-shell-*/chrome-headless-shell")), reverse=True)
    candidates = [
        os.environ.get("CHROME_BIN"),
        *shells,  # exits cleanly after one capture; desktop Chrome can linger, see chrome()
        shutil.which("chrome-headless-shell"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Chromium.app/Contents/MacOS/Chromium",
        shutil.which("google-chrome"), shutil.which("chromium"), shutil.which("chromium-browser"),
        *sorted(glob.glob(os.path.expanduser(
            "~/Library/Caches/ms-playwright/chromium-*/chrome-mac*/Chromium.app/Contents/MacOS/Chromium")), reverse=True),
    ]
    for c in candidates:
        if c and os.path.exists(c):
            return c
    sys.exit("Chrome/Chromium not found. Set CHROME_BIN, or run with --skip-capture after saving the PNGs.")


def serve() -> tuple[http.server.ThreadingHTTPServer, str]:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]

    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *args) -> None:  # noqa: D401 - silence request log
            pass

    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", port), functools.partial(Quiet, directory=str(LANDING)))
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd, f"http://127.0.0.1:{port}/deck/"


def chrome(binary: str, profile: str, out: Path, *args: str) -> None:
    """Run one headless capture that writes `out`.

    Desktop Chrome sometimes keeps running after it has written the file, so the output is watched:
    once it exists and has stopped growing, the browser is stopped.
    """
    out.unlink(missing_ok=True)
    cmd = [binary, "--hide-scrollbars", "--no-first-run", "--no-default-browser-check", "--force-device-scale-factor=1",
           f"--user-data-dir={profile}", f"--window-size={W},{H}", "--virtual-time-budget=8000", *args]
    if "headless-shell" not in os.path.basename(binary):
        cmd[1:1] = ["--headless=new", "--disable-gpu"]
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    deadline, last, stable = time.time() + 120, -1, 0
    while proc.poll() is None and time.time() < deadline:
        time.sleep(0.25)
        size = out.stat().st_size if out.exists() else -1
        stable = stable + 1 if size > 0 and size == last else 0
        last = size
        if stable >= 8:  # unchanged for 2 s
            break
    if proc.poll() is None:
        proc.terminate()
        try:
            proc.wait(10)
        except subprocess.TimeoutExpired:
            proc.kill()
    if not out.exists() or out.stat().st_size == 0:
        sys.exit(f"Chrome did not write {out}")


def capture(n_slides: int) -> None:
    binary = find_chrome()
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    for old in PNG_DIR.glob("slide-*.png"):
        old.unlink()
    httpd, base = serve()
    profile = tempfile.mkdtemp(prefix="deck-export-")
    try:
        for i in range(1, n_slides + 1):
            out = PNG_DIR / f"slide-{i:02d}.png"
            chrome(binary, profile, out, f"--screenshot={out}", f"{base}?export={i}")
            print(f"  png  {out.relative_to(ROOT)}")
        PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
        chrome(binary, profile, PDF_OUT, f"--print-to-pdf={PDF_OUT}", "--no-pdf-header-footer", f"{base}?print")
        print(f"  pdf  {PDF_OUT.relative_to(ROOT)}")
    finally:
        httpd.shutdown()
        shutil.rmtree(profile, ignore_errors=True)


def normalise_png(path: Path) -> None:
    """Guarantee an exact 1920x1080 RGB image whatever the browser produced."""
    with Image.open(path) as im:
        if im.size == (W, H) and im.mode == "RGB":
            return
        im = im.convert("RGB")
        if im.size != (W, H):
            scale = max(W / im.width, H / im.height)
            im = im.resize((round(im.width * scale), round(im.height * scale)), Image.LANCZOS).crop((0, 0, W, H))
        im.save(path, optimize=True)


def build_pptx(slides: list[dict[str, str]]) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * EMU_PER_PX), Emu(H * EMU_PER_PX)
    prs.core_properties.title = "AI Portfolio Advisor: capstone deck"
    prs.core_properties.subject = "A forecast is a range"
    blank = prs.slide_layouts[6]
    for i, meta in enumerate(slides, start=1):
        png = PNG_DIR / f"slide-{i:02d}.png"
        if not png.exists():
            sys.exit(f"Missing {png}. Run without --skip-capture first.")
        normalise_png(png)
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = f"Slide {i}: {meta['title']}"
        pic._element.nvPicPr.cNvPr.set("descr", f"{meta['title']} {meta['notes']}"[:2000])
        notes = [meta["title"], f"Section: {meta['section']}", "", meta["notes"]]
        if i == len(slides):
            notes += ["", f"PowerPoint cannot carry the HTML animations and builds; the animated version is at {LIVE_URL}"]
        slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    PPTX_OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_OUT)
    print(f"  pptx {PPTX_OUT.relative_to(ROOT)}")


def pdf_pages(path: Path) -> int:
    return len(re.findall(rb"/Type\s*/Page(?![s\w])", path.read_bytes()))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--skip-capture", action="store_true", help="reuse PNGs in landing/deck/export/ and the existing PDF")
    args = ap.parse_args()

    slides = read_slides()
    print(f"{len(slides)} slides in landing/deck/index.html")
    if not args.skip_capture:
        capture(len(slides))
    build_pptx(slides)

    check = Presentation(PPTX_OUT)
    sizes = {Image.open(PNG_DIR / f"slide-{i:02d}.png").size for i in range(1, len(slides) + 1)}
    print(f"PPTX: {len(check.slides)} slides, {check.slide_width / check.slide_height:.4f} aspect, images {sorted(sizes)}, {PPTX_OUT.stat().st_size / 1e6:.2f} MB")
    if PDF_OUT.exists():
        print(f"PDF:  {pdf_pages(PDF_OUT)} pages, {PDF_OUT.stat().st_size / 1e6:.2f} MB")
    else:
        print("PDF:  not written (capture was skipped and no PDF exists)")


if __name__ == "__main__":
    main()
