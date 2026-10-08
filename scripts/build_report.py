"""Build the capstone written report (Word + Markdown + figures) from the saved metrics.

    python scripts/build_report.py            # figures, docs/REPORT.md, docs/Capstone_Report.docx (+ PDF if LibreOffice is installed)
    python scripts/build_report.py --no-pdf   # skip the PDF and the page-numbered table of contents

Inputs  : reports/metrics.json, reports/cv_results.csv  (both written by `python -m src.model`)
Outputs : docs/figures/*.png, docs/REPORT.md, docs/Capstone_Report.docx, docs/Capstone_Report.pdf

Every number in the report is read from the two input files or computed from the
ROI assumptions in `ROI` below. The narrative interprets the committed metrics;
`check_narrative()` warns if a retrain changes a result the prose depends on.
The report content is defined once (see `build_content`) and rendered twice,
so the Word file and the Markdown file cannot drift apart.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
FIGS = DOCS / "figures"
DOCX_PATH = DOCS / "Capstone_Report.docx"
PDF_PATH = DOCS / "Capstone_Report.pdf"
MD_PATH = DOCS / "REPORT.md"

M = json.loads((ROOT / "reports" / "metrics.json").read_text())
CV = pd.read_csv(ROOT / "reports" / "cv_results.csv")
D, CFG, SEL, HOLD, BT = M["data"], M["config"], M["selected"], M["holdout"], M["backtest"]
LG, NV, RG = HOLD["lightgbm"], HOLD["naive"], HOLD["ridge"]
CAND = {c["model"]: c for c in M["candidates"]}
NAIVE = CAND["Naive (historical quantiles)"]
RIDGE = CAND["Ridge (scaled features)"]
NOTEBOOK = CAND["LightGBM quantile, notebook settings"]
TUNED = CAND["LightGBM quantile, tuned"]
TRIALS = pd.DataFrame(M["trials"])
BEST_TRIAL = int(TRIALS.iloc[0]["trial"])
FOLDS = M["validation"]["folds"]
GAP = CFG["HORIZON"] + CFG["EMBARGO"]
BH, BC = BT["holdout"], BT["cross_validation_folds"]
ERR = M["error_analysis"]

APP_URL = "https://fn.iimbg.com"
LANDING_URL = "https://fin.iimbg.com"
REPO_URL = "https://github.com/build-with-bala/ai-portfolio-advisor"

COURSE = "AI and ML for Digital Business Managers  ·  Applied AI & Machine Learning Capstone"
SUBTITLE = "Calibrated 21-day return ranges for NSE large caps, explained and sized to an investor's risk profile"
GROUP = "Group 5"
TEAM = [("B Gokul Vaigundh", "DBM/1015/03"), ("Balaji G", "DBM/1016/03"), ("Shashaanth NC", "DBM/1070/03")]
DISCLAIMER = ("AI Portfolio Advisor is a research and education tool. It is not investment advice, its authors are not registered "
              "investment advisers, and it places no trades.")

# Adviser-productivity case. Every input is an ASSUMPTION, not a measurement.
ROI = dict(clients=200, reviews_per_year=4, minutes_before=45, minutes_after=10, adviser_rate_inr=1500, annual_cost_inr=120_000)


def roi_case(clients: int = ROI["clients"], minutes_saved: float = ROI["minutes_before"] - ROI["minutes_after"],
             rate: float = ROI["adviser_rate_inr"], cost: float = ROI["annual_cost_inr"]) -> dict[str, float]:
    hours = clients * ROI["reviews_per_year"] * minutes_saved / 60
    value = hours * rate
    return {"hours": hours, "value": value, "net": value - cost, "roi": (value - cost) / cost}


BASE = roi_case()

# ------------------------------------------------------------------ formatting
def pct(x: float, d: int = 1) -> str:
    return f"{x * 100:.{d}f}%"


def spct(x: float, d: int = 1) -> str:
    return f"{x * 100:+.{d}f}%"


def f5(x: float) -> str:
    return f"{x:.5f}"


def f4(x: float) -> str:
    return f"{x:.4f}"


def sic(x: float) -> str:
    return f"{x:+.3f}"


def inr(x: float) -> str:
    return f"Rs {x:,.0f}"


def lakh(x: float, d: int = 1) -> str:
    return f"Rs {x / 1e5:.{d}f} lakh"


def params_text(p: dict) -> str:
    return ", ".join(f"{k}={v:g}" for k, v in p.items())


# --------------------------------------------------------------------- figures
INK, INDIGO, MARIGOLD, MUTED = "#141936", "#3B46C4", "#D98E04", "#5B5F77"
BRAND_MARIGOLD = "#F2A51A"  # product colour; too light for data marks on white, used for accents only
NEUTRAL, GRID, AXIS = "#A3A6B8", "#E8E9F0", "#C9CBD6"
TINT_INDIGO, TINT_MARIGOLD, TINT_GREY = "#ECEEFB", "#FDF3DC", "#F4F5F9"

plt.rcParams.update({
    "font.family": ["Arial", "Helvetica", "DejaVu Sans"],
    "font.size": 9, "text.color": INK, "axes.labelcolor": MUTED, "axes.edgecolor": AXIS,
    "xtick.color": MUTED, "ytick.color": MUTED, "xtick.labelsize": 8, "ytick.labelsize": 8,
    "axes.titlesize": 10.5, "axes.titleweight": "bold", "axes.titlelocation": "left", "axes.titlepad": 10,
    "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white",
    "axes.linewidth": 0.8, "legend.frameon": False, "legend.fontsize": 8,
})


def _style(ax, grid: str = "y") -> None:
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.tick_params(length=0)
    ax.set_axisbelow(True)
    if grid:
        ax.grid(axis=grid, color=GRID, linewidth=0.8)


def _save(fig, name: str) -> Path:
    FIGS.mkdir(parents=True, exist_ok=True)
    path = FIGS / f"{name}.png"
    fig.savefig(path, dpi=200, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)
    return path


def fig_architecture() -> Path:
    fig, ax = plt.subplots(figsize=(8.6, 4.9))
    ax.set_xlim(0, 100), ax.set_ylim(0, 64), ax.axis("off")

    def box(x, y, w, h, title, body, fill, edge):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=1.2", facecolor=fill, edgecolor=edge, linewidth=1.1))
        ax.text(x + w / 2, y + h - 2.6, title, ha="center", va="top", fontsize=8.6, fontweight="bold", color=INK)
        ax.text(x + w / 2, y + h - 6.4, body, ha="center", va="top", fontsize=7.4, color=MUTED, linespacing=1.35)

    def arrow(a, b, rad=0.0):
        ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=9, linewidth=1.0, color=MUTED,
                                     connectionstyle=f"arc3,rad={rad}", shrinkA=1.5, shrinkB=1.5))

    for x, label in [(1, "SOURCES"), (26, "TRAINING"), (53, "ARTEFACTS"), (75, "PRODUCT AND REPORTING")]:
        ax.text(x, 62.5, label, fontsize=7.4, fontweight="bold", color=MUTED)
    ax.plot([1, 99], [60.3, 60.3], color=GRID, linewidth=0.8)

    box(1, 40, 20, 17, "Yahoo Finance", "daily adjusted prices,\ntoday's fundamentals;\nkept in data/snapshot/", "white", AXIS)
    box(1, 2, 20, 17, "Live quotes", "optional: Upstox,\nZerodha or Yahoo,\nthrough live.py", "white", AXIS)
    box(26, 40, 22, 17, "src/data.py", "19 candidate features\nfractional differencing\nfeature selection\npurged time splits", "white", INDIGO)
    box(26, 14, 22, 20, "src/model.py", "25-setting random search\n5-fold walk-forward CV\n3 quantile models\nhold-out test, SHAP,\nbacktest", "white", INDIGO)
    box(53, 42, 17, 15, "reports/", "metrics.json\ncv_results.csv", TINT_MARIGOLD, MARIGOLD)
    box(53, 16, 17, 16, "data/artifacts.pkl", "models, forecasts,\nprices, fundamentals", TINT_MARIGOLD, MARIGOLD)
    box(75, 44, 24, 13, "scripts/build_report.py", "figures and this report", TINT_GREY, AXIS)
    box(75, 24, 24, 15, "portfolio_advisor/core.py", "four pillar scores, profile\ngates, HRP position sizing", TINT_INDIGO, INDIGO)
    box(75, 2, 24, 17, "app.py  (Streamlit)", "seven tabs and analyst chat\nlive at fn.iimbg.com\napi_app.py: JSON API", TINT_INDIGO, INDIGO)

    arrow((21, 48.5), (26, 48.5))
    arrow((37, 40), (37, 34))
    arrow((48, 29), (53, 47), rad=-0.12)
    arrow((48, 23), (53, 23))
    arrow((70, 50), (75, 50))
    arrow((70, 25), (75, 30))
    arrow((87, 24), (87, 19))
    arrow((21, 8), (75, 8))
    ax.text(48, 9.2, "intraday overlay only; never retrains the model", fontsize=6.8, color=MUTED, ha="center")
    return _save(fig, "fig_architecture")


def fig_roi() -> Path:
    fig, ax = plt.subplots(figsize=(7.2, 3.7))
    clients = np.arange(0, 401, 10)
    base_saved = ROI["minutes_before"] - ROI["minutes_after"]
    for saved in (15, 25, 35, 45):
        net = np.array([roi_case(c, saved)["net"] for c in clients]) / 1e5
        is_base = saved == base_saved
        ax.plot(clients, net, color=INDIGO if is_base else NEUTRAL, linewidth=2.2 if is_base else 1.4, solid_capstyle="round", zorder=3 if is_base else 2)
        ax.text(404, net[-1], f"{saved} min saved" + (" (base case)" if is_base else ""), va="center", fontsize=8, color=INK if is_base else MUTED)
    ax.axhline(0, color=INK, linewidth=0.9)
    ax.text(2, 0.35, "break-even", fontsize=7.5, color=MUTED, va="bottom")
    ax.scatter([ROI["clients"]], [BASE["net"] / 1e5], s=62, color=INDIGO, edgecolor="white", linewidth=2, zorder=5)
    ax.annotate(f"Base case: {ROI['clients']} clients\nnet benefit {lakh(BASE['net'])} a year", (ROI["clients"], BASE["net"] / 1e5), xytext=(-12, 24),
                textcoords="offset points", ha="right", fontsize=8, color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, linewidth=0.7))
    ax.set_xlim(0, 400), ax.set_xlabel("Clients reviewed four times a year")
    ax.set_ylabel("Net annual benefit (Rs lakh)")
    ax.set_title("Net benefit grows with the client book; a small practice barely breaks even")
    _style(ax)
    return _save(fig, "fig_roi_sensitivity")


def fig_pinball() -> Path:
    fig, ax = plt.subplots(figsize=(6.6, 3.3))
    err = np.linspace(-0.10, 0.10, 201)
    for q, color, label in [(0.1, MARIGOLD, "P10 model (q = 0.1)"), (0.5, INK, "P50 model (q = 0.5)"), (0.9, INDIGO, "P90 model (q = 0.9)")]:
        ax.plot(err * 100, np.maximum(q * err, (q - 1) * err) * 100, color=color, linewidth=2, label=label, solid_capstyle="round")
    ax.set_xticks(range(-10, 11, 2))
    ax.set_xlabel("Actual return minus forecast (percentage points)")
    ax.set_ylabel("Pinball loss (percentage points)")
    ax.set_title("Pinball loss is a tilted absolute error: the tilt sets the quantile")
    ax.legend(loc="upper center", ncol=1, handlelength=1.6)
    _style(ax)
    return _save(fig, "fig_pinball_loss")


def fig_selection() -> Path:
    table = pd.DataFrame(M["feature_selection"]["table"]).sort_values("relevance")
    fig, ax = plt.subplots(figsize=(7.2, 5.0))
    y = np.arange(len(table))
    ax.barh(y, table["relevance"], height=0.56, color=[INDIGO if s else NEUTRAL for s in table["selected"]], zorder=2)
    chosen = (table["selected"] & (table["redundancy"] > 0)).to_numpy()
    ax.scatter(table["redundancy"][chosen], y[chosen], s=46, color=MARIGOLD, edgecolor="white", linewidth=1.6, zorder=4)
    ax.set_yticks(y, table["feature"])
    ax.tick_params(axis="y", labelsize=8, labelcolor=INK)
    ax.set_xlabel("Mutual information (nats)")
    ax.set_title("Relevance to the label, and redundancy with features already chosen")
    handles = [Rectangle((0, 0), 1, 1, color=INDIGO), Rectangle((0, 0), 1, 1, color=NEUTRAL),
               plt.Line2D([], [], marker="o", linestyle="", color=MARIGOLD, markersize=6.5)]
    ax.legend(handles, ["Relevance, selected (10)", "Relevance, not selected (9)", "Redundancy when it was selected"], loc="lower right")
    _style(ax, grid="x")
    return _save(fig, "fig_feature_selection")


def fig_folds() -> Path:
    fig, ax = plt.subplots(figsize=(7.4, 3.4))
    num = lambda text: float(mdates.date2num(pd.Timestamp(text)))  # noqa: E731
    first, gap_days = num(D["first_date"]), GAP * 365.25 / 252
    rows = [(f"Fold {f['fold']}", first, num(f["val_start"]), num(f["val_end"])) for f in FOLDS]
    rows.append(("Final model\nand hold-out", first, num(D["holdout_start"]), num(D["holdout_end_scored"])))
    for i, (label, start, val_start, val_end) in enumerate(rows):
        yy = len(rows) - 1 - i
        ax.barh(yy, val_start - gap_days - start, left=start, height=0.5, color="#C7CBEE", zorder=2)
        ax.barh(yy, gap_days, left=val_start - gap_days, height=0.5, color=MARIGOLD, zorder=3)
        ax.barh(yy, val_end - val_start, left=val_start, height=0.5, color=INDIGO if i < len(FOLDS) else INK, zorder=2)
    ax.set_yticks(range(len(rows)), [r[0] for r in rows][::-1])
    ax.tick_params(axis="y", labelsize=8, labelcolor=INK)
    ax.xaxis_date()
    ax.set_xlim(first - 40, num(D["last_date"]) + 40)
    handles = [Rectangle((0, 0), 1, 1, color=c) for c in ("#C7CBEE", MARIGOLD, INDIGO, INK)]
    ax.legend(handles, ["Training data", f"Purged gap ({GAP} trading days)", "Validation block", "Hold-out test"], loc="upper center",
              bbox_to_anchor=(0.5, -0.12), ncol=4, columnspacing=1.4, handlelength=1.2)
    ax.set_title("Each fold trains on the past, skips a purged gap, then validates on the next block")
    _style(ax, grid="x")
    return _save(fig, "fig_walk_forward_folds")


def _lgbm_fold(trial: int, column: str) -> np.ndarray:
    rows = CV[(CV["model"] == "LightGBM quantile") & (CV["trial"] == trial)].sort_values("fold")
    return rows[column].to_numpy()


def _naive_fold(column: str) -> np.ndarray:
    return CV[CV["model"].str.startswith("Naive")].sort_values("fold")[column].to_numpy()


def fig_cv_candidates() -> Path:
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    folds = np.arange(1, len(FOLDS) + 1)
    series = [("LightGBM, notebook settings", _lgbm_fold(0, "pinball"), MARIGOLD, 2), ("Naive baseline", _naive_fold("pinball"), NEUTRAL, 3),
              ("LightGBM, tuned", _lgbm_fold(BEST_TRIAL, "pinball"), INDIGO, 4)]
    for label, values, color, z in series:
        ax.plot(folds, values, color=color, linewidth=2, marker="o", markersize=6.5, markeredgecolor="white", markeredgewidth=1.6, label=label, zorder=z)
    ax.annotate("notebook settings\nover-fit the early folds", (2, _lgbm_fold(0, "pinball")[1]), xytext=(18, 2), textcoords="offset points", fontsize=7.8, color=MUTED, va="center")
    ax.annotate("tuned model and naive\nbaseline are almost identical", (4, _naive_fold("pinball")[3]), xytext=(0, -26), textcoords="offset points", fontsize=7.8, color=MUTED, ha="center")
    ax.set_xticks(folds, [f"Fold {f['fold']}\n{f['val_start'][:4]}-{f['val_end'][2:4]}" for f in FOLDS])
    ax.set_xlim(0.7, 5.3), ax.set_ylim(0.012, 0.037)
    ax.set_ylabel("Mean pinball loss (lower is better)")
    ax.set_title("The period matters more than the model: loss by validation fold")
    ax.legend(loc="upper right")
    _style(ax)
    return _save(fig, "fig_cv_candidates")


def fig_trials() -> Path:
    fig, ax = plt.subplots(figsize=(7.4, 3.7))
    order = TRIALS.reset_index(drop=True)
    for rank, row in order.iterrows():
        trial = int(row["trial"])
        values = _lgbm_fold(trial, "pinball")
        color = INDIGO if trial == BEST_TRIAL else MARIGOLD if trial == 0 else NEUTRAL
        ax.plot([rank + 1, rank + 1], [values.min(), values.max()], color=color, linewidth=1.3, alpha=0.9, zorder=2)
        ax.scatter([rank + 1], [row["mean"]], s=44, color=color, edgecolor="white", linewidth=1.5, zorder=3)
    ax.axhline(NAIVE["pinball_mean"], color=INK, linewidth=0.9, zorder=1)
    ax.text(len(order) + 0.9, NAIVE["pinball_mean"], "naive\nbaseline", va="center", fontsize=7.6, color=MUTED)
    best = order.iloc[0]
    ax.annotate(f"chosen: {int(best['num_leaves'])} leaves,\n{int(best['n_estimators'])} trees", (1, best["mean"]), xytext=(14, -36), textcoords="offset points", fontsize=7.8,
                color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, linewidth=0.7), bbox=dict(facecolor="white", edgecolor="none", pad=2), zorder=6)
    nb_rank = int(order.index[order["trial"] == 0][0]) + 1
    nb = order.iloc[nb_rank - 1]
    ax.annotate(f"notebook: {int(nb['num_leaves'])} leaves,\n{int(nb['n_estimators'])} trees", (nb_rank, nb["mean"]), xytext=(-70, -52), textcoords="offset points", fontsize=7.8,
                color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, linewidth=0.7), bbox=dict(facecolor="white", edgecolor="none", pad=2), zorder=6)
    ax.set_xlim(0, len(order) + 2.6), ax.set_xticks([1, 5, 10, 15, 20, 25])
    ax.set_xlabel("Hyperparameter setting, ranked by mean pinball loss")
    ax.set_ylabel("Pinball loss (dot = mean of 5 folds,\nline = best to worst fold)")
    ax.set_title("25 settings: differences between settings are small next to the spread across folds")
    _style(ax)
    return _save(fig, "fig_hyperparameter_trials")


def fig_coverage() -> Path:
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    x = np.arange(1, len(FOLDS) + 2)
    hold_x = x[-1]
    series = [("LightGBM, notebook settings", _lgbm_fold(0, "coverage"), None, MARIGOLD), ("Naive baseline", _naive_fold("coverage"), NV["coverage"], NEUTRAL),
              ("LightGBM, tuned", _lgbm_fold(BEST_TRIAL, "coverage"), LG["coverage"], INDIGO)]
    for label, cv_values, hold, color in series:
        ax.plot(x[:-1], cv_values * 100, color=color, linewidth=2, marker="o", markersize=6.5, markeredgecolor="white", markeredgewidth=1.6, label=label)
        if hold is not None:
            ax.scatter([hold_x], [hold * 100], s=62, color=color, edgecolor="white", linewidth=1.8, zorder=4)
    ax.axhline(HOLD["target_coverage"] * 100, color=INK, linewidth=0.9)
    ax.text(0.62, HOLD["target_coverage"] * 100 + 1.2, "80% target", fontsize=7.8, color=MUTED)
    ax.axvline(5.5, color=AXIS, linewidth=0.8)
    ax.annotate(f"{pct(LG['coverage'])}", (hold_x, LG["coverage"] * 100), xytext=(9, -4), textcoords="offset points", fontsize=8.2, color=INK, va="center")
    ax.annotate(f"{pct(NV['coverage'])}", (hold_x, NV["coverage"] * 100), xytext=(9, 4), textcoords="offset points", fontsize=8.2, color=MUTED, va="center")
    ax.set_xticks(x, [f"Fold {f['fold']}" for f in FOLDS] + ["Hold-out\n2025-26"])
    ax.set_xlim(0.55, 6.75), ax.set_ylim(40, 100)
    ax.set_ylabel("Outcomes inside the P10 to P90 band (%)")
    ax.set_title("Calibration: how often the real return landed inside the forecast band")
    ax.legend(loc="lower right")
    _style(ax)
    return _save(fig, "fig_coverage_calibration")


def fig_shap() -> Path:
    shap = pd.DataFrame(M["shap"]).iloc[::-1]
    fig, ax = plt.subplots(figsize=(7.0, 3.6))
    y = np.arange(len(shap))
    values = shap["mean_abs_shap"] * 100
    ax.barh(y, values, height=0.56, color=[INDIGO if d >= 0 else MARIGOLD for d in shap["direction"]], zorder=2)
    for yy, v in zip(y, values):
        ax.text(v + 0.006, yy, f"{v:.3f}", va="center", fontsize=7.8, color=INK)
    ax.set_yticks(y, shap["feature"])
    ax.tick_params(axis="y", labelsize=8, labelcolor=INK)
    ax.set_xlim(0, values.max() * 1.16)
    ax.set_xlabel("Mean absolute SHAP value (percentage points of 21-day return)")
    ax.set_title("Five features carry the median forecast; the other five barely move it")
    handles = [Rectangle((0, 0), 1, 1, color=INDIGO), Rectangle((0, 0), 1, 1, color=MARIGOLD)]
    ax.legend(handles, ["Higher feature value raises the forecast", "Higher feature value lowers the forecast"], loc="lower right")
    _style(ax, grid="x")
    return _save(fig, "fig_shap_importance")


def fig_error_regime() -> Path:
    groups = ERR["volatility_regime"] + ERR["market_direction"]
    labels = ["Calm", "Normal", "Turbulent", "Market\nfell", "Market\nrose"]
    fig, axes = plt.subplots(1, 3, figsize=(8.4, 3.2))
    specs = [("mae", "Median-forecast error (MAE, pp)", 100, "{:.1f}"), ("coverage", "P10 to P90 coverage (%); line = 80% target", 100, "{:.0f}"), ("ic_mean", "Rank IC (stock ordering)", 1, "{:+.2f}")]
    for ax, (key, title, scale, fmt) in zip(axes, specs):
        values = np.array([g[key] for g in groups]) * scale
        ax.bar(range(5), values, width=0.5, color=INDIGO, zorder=2)
        for i, v in enumerate(values):
            pad = abs(values).max() * 0.035
            top = max(v, 80) if key == "coverage" else v
            ax.text(i, top + (pad if v >= 0 else -pad), fmt.format(v), ha="center", va="bottom" if v >= 0 else "top", fontsize=7.6, color=INK)
        ax.axvline(2.5, color=AXIS, linewidth=0.8)
        ax.set_xticks(range(5), labels, fontsize=7.2)
        ax.set_title(title, fontsize=8.6, fontweight="bold")
        if key == "coverage":
            ax.axhline(80, color=INK, linewidth=0.9)
            ax.set_ylim(0, 100)
        if key == "ic_mean":
            ax.axhline(0, color=INK, linewidth=0.9)
            ax.set_ylim(-0.13, 0.18)
        if key == "mae":
            ax.set_ylim(0, 8)
        _style(ax)
    fig.suptitle("Hold-out forecast quality by volatility regime and by market direction", x=0.01, ha="left", fontsize=10.5, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    return _save(fig, "fig_error_regime")


def fig_error_sector() -> Path:
    sectors = pd.DataFrame(ERR["sector"]).sort_values("coverage")
    fig, axes = plt.subplots(1, 2, figsize=(8.4, 3.9), sharey=True)
    y = np.arange(len(sectors))
    ax = axes[0]
    ax.hlines(y, 80, sectors["coverage"] * 100, color=NEUTRAL, linewidth=1.2, zorder=2)
    ax.scatter(sectors["coverage"] * 100, y, s=48, color=INDIGO, edgecolor="white", linewidth=1.6, zorder=3)
    ax.axvline(80, color=INK, linewidth=0.9)
    ax.set_xlim(68, 92), ax.set_ylim(-0.7, len(sectors) - 0.1)
    ax.set_yticks(y, sectors["group"])
    ax.tick_params(axis="y", labelsize=7.8, labelcolor=INK)
    ax.set_title("P10 to P90 coverage (%); line = 80% target", fontsize=8.6, fontweight="bold")
    _style(ax, grid="x")
    ax = axes[1]
    ax.barh(y, sectors["mae"] * 100, height=0.5, color=INDIGO, zorder=2)
    for yy, v in zip(y, sectors["mae"] * 100):
        ax.text(v + 0.1, yy, f"{v:.1f}", va="center", fontsize=7.6, color=INK)
    ax.set_xlim(0, 8.6)
    ax.set_title("Median-forecast error (MAE, pp)", fontsize=8.6, fontweight="bold")
    _style(ax, grid="x")
    fig.suptitle("Hold-out forecast quality by sector (five stocks each)", x=0.01, ha="left", fontsize=10.5, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return _save(fig, "fig_error_sector")


def fig_backtest() -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(8.4, 3.4))
    panels = [(BC["first_run"], f"Cross-validation years, {BC['first_run']['start'][:4]} to {BC['first_run']['end'][:4]}"),
              (BH["first_run"], f"Hold-out, {BH['first_run']['start'][:7]} to {BH['first_run']['end'][:7]}")]
    for ax, (run, title) in zip(axes, panels):
        strategy = np.r_[1.0, run["curve"]["strategy"]]
        benchmark = np.r_[1.0, run["curve"]["benchmark"]]
        x = np.arange(len(strategy))
        ax.plot(x, benchmark, color=NEUTRAL, linewidth=2, label="Benchmark: all 60 stocks, equal weight", solid_capstyle="round")
        ax.plot(x, strategy, color=INDIGO, linewidth=2, label="Strategy: top 5 median forecasts", solid_capstyle="round")
        ax.scatter([x[-1], x[-1]], [benchmark[-1], strategy[-1]], s=40, color=[NEUTRAL, INDIGO], edgecolor="white", linewidth=1.5, zorder=4)
        offset = 7 if strategy[-1] >= benchmark[-1] else -7
        ax.annotate(f"{strategy[-1]:.2f}", (x[-1], strategy[-1]), xytext=(6, offset), textcoords="offset points", fontsize=8, color=INK, va="center")
        ax.annotate(f"{benchmark[-1]:.2f}", (x[-1], benchmark[-1]), xytext=(6, -offset), textcoords="offset points", fontsize=8, color=MUTED, va="center")
        ax.axhline(1, color=AXIS, linewidth=0.8)
        ax.set_xlim(0, len(x) * 1.1)
        ax.set_xlabel("Rebalance number (one every 21 trading days)")
        ax.set_title(title, fontsize=8.6, fontweight="bold")
        _style(ax)
    axes[0].set_ylabel("Value of 1 rupee invested")
    axes[0].legend(loc="upper left")
    fig.suptitle("Backtest after 20 bps costs: no reliable edge over simply holding every stock", x=0.01, ha="left", fontsize=10.5, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return _save(fig, "fig_backtest")


FIGURE_BUILDERS = {
    "architecture": fig_architecture, "roi": fig_roi, "pinball": fig_pinball, "selection": fig_selection, "folds": fig_folds,
    "cv": fig_cv_candidates, "trials": fig_trials, "coverage": fig_coverage, "shap": fig_shap, "regime": fig_error_regime,
    "sector": fig_error_sector, "backtest": fig_backtest,
}


def build_figures() -> dict[str, Path]:
    return {key: builder() for key, builder in FIGURE_BUILDERS.items()}


# --------------------------------------------------------------- content model
# The report is a flat list of blocks. Inline markup: **bold**, *italic*, `code`.
# Cross-references: [[fig:key]] and [[tab:key]] become "Figure 3" / "Table 5".
BLOCKS: list[dict] = []


def _add(kind: str, **fields) -> None:
    BLOCKS.append({"kind": kind, **fields})


def h1(text: str, page_break: bool = False) -> None:
    _add("h1", text=text, page_break=page_break)


def h2(text: str) -> None:
    _add("h2", text=text)


def p(text: str) -> None:
    _add("p", text=" ".join(text.split()))


def bullets(items: list[str]) -> None:
    _add("bullets", items=[" ".join(i.split()) for i in items])


def numbered(items: list[str]) -> None:
    _add("numbered", items=[" ".join(i.split()) for i in items])


def table(key: str, caption: str, header: list[str], rows: list[list], widths: list[float], align: str | None = None, size: float = 8.5) -> None:
    _add("table", key=key, caption=caption, header=header, rows=[[str(c) for c in r] for r in rows], widths=widths,
         align=align or "l" * len(header), size=size)


def figure(key: str, caption: str, width_cm: float = 15.6) -> None:
    _add("figure", key=key, caption=" ".join(caption.split()), width_cm=width_cm)


def callout(title: str, text: str) -> None:
    _add("callout", title=title, text=" ".join(text.split()))


def eq(text: str) -> None:
    _add("eq", text=text)


def code(lines: list[str]) -> None:
    _add("code", lines=lines)


def qa(question: str, answer: str) -> None:
    _add("qa", q=" ".join(question.split()), a=" ".join(answer.split()))


def universe_rows() -> list[list[str]]:
    """Sector universe from src/data.py (the pipeline's own definition)."""
    sys.path.insert(0, str(ROOT))
    sys.path.insert(0, str(ROOT / "src"))
    from src import data as pipeline  # noqa: PLC0415 - imported late so figures can be built without the ML stack

    return [[sector, ", ".join(t.replace(".NS", "") for t in tickers)] for sector, tickers in pipeline.SECTOR_UNIVERSE.items()]


FEATURES = [
    ("RSI_14", "Relative Strength Index over 14 days: 100 − 100 / (1 + average gain / average loss), with exponential (Wilder) smoothing", "Whether recent gains have outweighed recent losses (0 to 100)"),
    ("MACD_signal", "MACD histogram (12-day minus 26-day exponential average, less its own 9-day average), divided by price", "Whether the short-term trend is accelerating or fading"),
    ("Boll_BW", "Bollinger band width: 4 × 20-day standard deviation of price ÷ 20-day average price", "How wide the recent trading range is"),
    ("ATR_14", "Average True Range over 14 days ÷ price", "Typical daily price range, including overnight gaps"),
    ("Stoch_K", "Stochastic %K: 100 × (close − 14-day low) ÷ (14-day high − 14-day low)", "Where today's close sits inside the 14-day range"),
    ("OBV_delta", "5-day percentage change in on-balance volume (running sum of volume, signed by the day's price direction), clipped to ±5", "Whether volume is flowing into up days or down days"),
    ("Williams_R", "−100 × (14-day high − close) ÷ (14-day high − 14-day low)", "The same position-in-range reading as Stoch_K, shifted by 100"),
    ("ret_5", "Price change over the last 5 trading days", "One-week momentum or reversal"),
    ("ret_21", "Price change over the last 21 trading days", "One-month momentum"),
    ("vol_21", "Standard deviation of daily returns over 21 days", "Recent stock-level risk"),
    ("fracdiff_close", "Log price, fractionally differenced with the smallest order d that passes the ADF test", "The price trend with its long memory kept but its drift removed"),
    ("ret_63", "Price change over the last 63 trading days", "One-quarter momentum"),
    ("mom_126_21", "Price 21 days ago ÷ price 126 days ago − 1", "Six-month momentum that skips the most recent month"),
    ("vol_63", "Standard deviation of daily returns over 63 days", "Slower-moving stock-level risk"),
    ("dist_52w_high", "Close ÷ highest close of the last 252 days − 1", "How far the stock is below its one-year high"),
    ("volume_surge", "5-day average volume ÷ 60-day average volume − 1", "Unusual trading activity"),
    ("rel_ret_21", "ret_21 minus the average ret_21 of all stocks that day", "One-month performance relative to the market"),
    ("mkt_ret_21", "Average ret_21 across all stocks that day", "Direction of the whole market over the last month"),
    ("mkt_vol_21", "Average vol_21 across all stocks that day", "How turbulent the whole market is"),
]

PARAM_NOTES = {
    "n_estimators": "Number of trees added one after another. More trees fit the training data more closely.",
    "learning_rate": "Fraction of each new tree's correction that is kept. Small values make learning slower and smoother.",
    "num_leaves": "Maximum end-points (leaves) per tree. Four leaves means at most three yes/no questions per tree.",
    "min_child_samples": "Minimum training rows a leaf must contain. Larger values stop the tree isolating small, noisy groups.",
    "subsample": "Share of training rows drawn at random for each tree.",
    "colsample_bytree": "Share of the features offered to each tree.",
    "reg_lambda": "L2 penalty that shrinks leaf values towards zero.",
}

GLOSSARY = [
    ("ADF test", "Augmented Dickey-Fuller test. A statistical test whose null hypothesis is that a series has a unit root (it wanders without a fixed mean). A p-value below 0.05 is taken as evidence that the series is stationary."),
    ("Backtest", "Replaying a trading rule on past data to see what it would have earned. It is an estimate under assumptions, not a track record."),
    ("Basis point (bp)", "One hundredth of one per cent. The backtest charges 20 bps (0.20%) on the value traded."),
    ("Coverage", "Share of real outcomes that fall inside the forecast band. An 80% band is well calibrated when coverage is close to 80%."),
    ("Cross-validation", "Estimating out-of-sample accuracy by repeatedly training on one part of the data and scoring on another."),
    ("Embargo", "Extra days dropped after the purge so that slow-moving information cannot bridge training and validation data."),
    ("Fractional differencing", "Subtracting a weighted sum of past values, with a non-integer order d between 0 and 1, to remove a trend while keeping more of the series' memory than ordinary differencing."),
    ("Gradient boosting", "Building a model as a sum of many small decision trees, each trained to correct the errors left by the trees before it."),
    ("Hold-out set", "The most recent 20% of dates, set aside before any modelling decision and scored once at the end."),
    ("HRP", "Hierarchical Risk Parity. A position-sizing method that groups correlated stocks and splits capital so that each group contributes a similar amount of risk."),
    ("Hyperparameter", "A setting chosen before training (for example tree size) as opposed to a value learned from the data."),
    ("Leakage (look-ahead)", "Any path by which information from the future reaches the model during training, making test scores better than real use would be."),
    ("LightGBM", "An open-source gradient-boosting library from Microsoft that grows trees leaf by leaf and bins feature values for speed."),
    ("MAE / RMSE", "Mean absolute error and root mean squared error of the median forecast, in return units (0.058 = 5.8 percentage points)."),
    ("Mutual information", "A measure of how much knowing one variable reduces uncertainty about another. It is zero when they are independent and captures non-linear relationships."),
    ("P10 / P50 / P90", "The 10th, 50th and 90th percentile forecasts: pessimistic case, median and optimistic case of the 21-day return."),
    ("Pinball loss", "The loss function minimised by a quantile forecast: an absolute error whose penalty is tilted according to the quantile."),
    ("Purging", "Removing training rows whose labels overlap in time with the validation period."),
    ("Quantile regression", "Regression that predicts a chosen percentile of the outcome instead of its average."),
    ("Rank IC", "Rank information coefficient: the Spearman rank correlation, across stocks on one day, between forecasts and the returns that followed. Zero means no ordering skill."),
    ("Ridge regression", "Linear regression with an L2 penalty (alpha) that shrinks coefficients towards zero."),
    ("SHAP value", "The contribution of one feature to one prediction, measured from the average prediction, using Shapley values from co-operative game theory."),
    ("Stationary series", "A series whose statistical behaviour (mean, variance) does not change over time, which is what most learning algorithms implicitly assume."),
    ("Survivorship bias", "Overstated historical performance caused by studying only companies that are still large today."),
    ("Walk-forward validation", "Cross-validation for time series: always train on the past and validate on the block that follows."),
]


def build_content() -> None:  # noqa: PLR0915 - one long, linear document
    params = SEL["params"]
    hold_gain = 1 - LG["pinball"] / NV["pinball"]
    band_cut = 1 - LG["band_width"] / NV["band_width"]
    purge_rows = D["panel_rows"] - D["development_rows"] - D["holdout_rows_scored"] - D["holdout_rows_latest_unlabelled"]
    saved = ROI["minutes_before"] - ROI["minutes_after"]
    break_even = ROI["annual_cost_inr"] / (ROI["reviews_per_year"] * saved / 60 * ROI["adviser_rate_inr"])
    shap = M["shap"]
    regime = {g["group"]: g for g in ERR["volatility_regime"]}
    direction = {g["group"]: g for g in ERR["market_direction"]}
    sectors = sorted(ERR["sector"], key=lambda g: g["coverage"])
    signals = {s["signal"]: s for s in HOLD["signals"]}
    sel_table = M["feature_selection"]["table"]
    n_selected = len(M["feature_selection"]["selected"])
    n_candidates = len(M["feature_selection"]["candidates"])
    fh, fc = BH["first_run"], BC["first_run"]

    _add("title")
    _add("toc")

    # ------------------------------------------------------ executive summary
    h1("Executive summary", page_break=True)
    p(f"""**The problem.** Independent wealth advisers and self-directed investors in India usually choose stocks from tips, broker
      notes and single price targets. A single target hides the two things a client most needs to know: how wide the range of likely
      outcomes is, and whether the position fits the client's tolerance for risk. Screening {D['tickers']} large companies by hand before
      each client review is also slow.""")
    p(f"""**The product.** AI Portfolio Advisor is a web application, live at {APP_URL}. The user sets a risk profile (Conservative,
      Moderate or Aggressive) and an amount of capital. The application returns a ranked book of {D['tickers']} National Stock Exchange
      (NSE) large caps, each marked STRONG BUY, BUY, WATCH or PASS, with a pessimistic, median and optimistic return for the next
      {CFG['HORIZON']} trading days, the evidence behind the rating, the model's reasons for that stock, and a position size in whole
      shares. No programming or statistics is needed to use it.""")
    p(f"""**The business case.** For an adviser with {ROI['clients']} clients reviewed {ROI['reviews_per_year']} times a year, cutting
      screening time from {ROI['minutes_before']} to {ROI['minutes_after']} minutes per review frees {BASE['hours']:,.0f} hours a year.
      At Rs {ROI['adviser_rate_inr']:,} an hour that is worth {lakh(BASE['value'])}, against an assumed running cost of
      {lakh(ROI['annual_cost_inr'])}: a return of {pct(BASE['roi'], 0)} on cost. These inputs are assumptions, stated openly and tested
      in a sensitivity table; they are not measurements.""")
    p(f"""**The machine learning.** Three LightGBM quantile-regression models forecast the 10th, 50th and 90th percentile of each
      stock's {CFG['HORIZON']}-day return from {n_selected} price-and-volume features, selected from {n_candidates} candidates by a
      mutual-information method. Settings were chosen by a random search over {M['validation']['lightgbm_trials']} configurations,
      each scored by purged walk-forward cross-validation on {CFG['N_SPLITS']} folds ({M['validation']['lightgbm_fits']} model fits).
      The most recent 20% of dates were locked away as a hold-out test. SHAP values explain every forecast.""")
    p(f"""**What the evidence shows.** Tuning cut the pinball loss (the error measure for quantile forecasts) by
      {pct(SEL['pinball_improvement_vs_notebook'])} compared with the settings in our original notebook. On the hold-out period,
      {pct(LG['coverage'])} of real outcomes fell inside the model's P10 to P90 band against a target of
      {pct(HOLD['target_coverage'], 0)}: the risk ranges are honest. However, the tuned model only ties a naive baseline that ignores
      all features ({f5(TUNED['pinball_mean'])} against {f5(NAIVE['pinball_mean'])} in cross-validation), and a strategy that buys the
      five highest forecasts did **not** beat holding all {D['tickers']} stocks equally on the hold-out period
      ({spct(BH['strategy_annual_return_median'])} a year against {spct(BH['benchmark_annual_return_median'])}).""")
    p("""**What we therefore claim.** The product saves screening time, enforces a consistent risk discipline, and gives clients a
      calibrated range of outcomes with a plain-language explanation. It does not claim to pick winning stocks, and this report says so
      wherever the question arises. It is a research and education tool, not investment advice.""")
    callout("How to read this report", f"""Sections 1 to 3 are written for a business reader. Sections 4 to 15 document the pipeline and
      every modelling choice. Section 16 states the limits. Appendix A lists the questions an examiner could ask any team member, with
      answers. Every number comes from `reports/metrics.json` (generated {M['generated']}) or is labelled an assumption.""")

    # ------------------------------------------------------------ 1 problem
    h1("1. Business problem and users", page_break=True)
    h2("1.1 The problem")
    p("""A retail investor who asks "should I buy this stock?" is normally given a point answer: a target price, a star rating or a tip.
      Three things are missing from that answer. First, uncertainty: a forecast of +2% over a month means little unless the client also
      knows that the plausible range runs from, say, −7% to +11%. Second, fit: the same stock can be reasonable for a 30-year-old with a
      long horizon and wrong for a retiree who cannot tolerate a 30% fall. Third, reasons: an adviser who cannot explain why a name is
      on the list cannot defend it to a client or to a compliance reviewer.""")
    p("""For the adviser the problem is also one of time. Before a client review the adviser has to look at price trends, valuation
      ratios, recent volatility and sector exposure for dozens of companies, and then translate that into a short list and position
      sizes that respect the client's risk limits. Done by hand, with a charting site and a spreadsheet, this is slow and inconsistent:
      two clients with the same profile can receive different advice on different days.""")
    h2("1.2 Target persona")
    p(f"""The primary user is an **independent wealth adviser** or a small advisory practice serving a few hundred retail clients, who
      reviews each client's equity holdings a few times a year. This user is financially literate but is not a data scientist and will
      not read code. The secondary user is a **self-directed investor** who reviews a portfolio about once a month and wants a
      disciplined second opinion. The forecast horizon of {CFG['HORIZON']} trading days (about one calendar month) was chosen to match
      that monthly review cycle.""")
    h2("1.3 Scope")
    p(f"""The scope is deliberately narrow so that it can be validated properly: {D['tickers']} liquid NSE large caps, five from each of
      {D['sectors']} sectors, daily data, one forecast horizon, long-only positions, and no order execution. The product produces
      research output for a human to review. It never connects to a brokerage account to place trades.""")

    # ------------------------------------------------------------ 2 product
    h1("2. The product in use")
    p(f"""The application is a single Streamlit web page ({APP_URL}). Streamlit is a Python framework that turns a script into an
      interactive web application, so the whole product, interface included, is Python. A user works from left to right: set the profile
      in the sidebar, read the summary strip, then open the tabs for detail.""")
    h2("2.1 Profile intake")
    p("""The sidebar holds a short form. The user picks a risk posture (Conservative, Moderate or Aggressive), enters investable capital
      and the maximum number of holdings, and may adjust seven limits that the posture pre-fills: minimum forecast confidence, minimum
      expected 21-day return, maximum annual volatility, maximum trailing drawdown, maximum weight of any one position, minimum
      fundamental score, and sector or ticker exclusions. Pressing "Run profile analysis" recomputes everything. A non-technical user
      can ignore the sliders entirely: choosing the posture and typing the capital is enough.""")
    h2("2.2 Summary strip and automated insights")
    p("""Six headline cards sit above the tabs: number of positions, the portfolio's base-case 21-day return (model median), its bear
      case (10th percentile), trailing annual volatility, the share of capital invested, and the average fundamental score. Below them,
      an "Automated insights" panel writes two or three sentences about the strongest names in plain language.""")
    h2("2.3 The seven tabs")
    table("tabs", "What each tab of the application shows", ["Tab", "What the user sees", "Decision it supports"], [
        ["Overview", "The recommendation book: every eligible stock with its rating, overall score, P10 / P50 / P90 return, confidence and target shares; decision cards; an evidence ledger with technical and fundamental readings and risk flags; near misses; a CSV download.", "Which names to discuss with the client, and why."],
        ["Stock research", "For any one ticker: candlestick chart with moving averages, Bollinger bands, volume, RSI and MACD; indicator readings; fundamental snapshot; the stock's own SHAP drivers.", "Whether a single recommendation survives a closer look."],
        ["Portfolio lab", "Capital map (weights and cash), bear / base / bull portfolio scenario, sector exposure, and an illustrative staged order schedule in whole shares.", "How much to hold of each name."],
        ["Validation", "Hold-out error, directional accuracy, rank IC, band coverage, and the realised return that followed each signal.", "How far to trust the forecasts."],
        ["Model report", "The full evidence in this report: cross-validation candidates, all hyperparameter trials, fold dates, hold-out test, backtest, SHAP, feature selection, error analysis and leakage controls.", "Technical audit of the model."],
        ["India live", "Optional intraday quotes through an Upstox or Zerodha exchange feed (or a delayed public fallback), shown as a separate overlay.", "Whether today's price action contradicts the daily forecast."],
        ["Method & data", "The decision pipeline in five steps, data coverage, and missing-data rates.", "What the tool does and does not use."],
    ], [2.6, 8.6, 4.8])
    p("""A floating "Ask the analyst" chat button lets the user ask questions in ordinary language (for example, "which banks have the
      narrowest forecast range?"). It is a large language model that is given the current table of all stocks as context and told to
      answer only from it; it is enabled only when an API key is configured and is not part of the validated forecast.""")
    h2("2.4 Design choices that serve a non-technical user")
    p("""Four choices were made for usability. Ratings use four familiar words instead of scores. Every forecast is shown as a range, never
      a single number. When no stock passes the user's limits the application says so and shows the near misses, instead of quietly
      relaxing the limits. And every screen that could be mistaken for advice carries the statement that the output is research only and
      that no broker connection or automatic order execution is enabled.""")

    # ---------------------------------------------------------------- 3 ROI
    h1("3. Business value and return on investment")
    h2("3.1 Where the value comes from")
    p("""The product creates value in three ways, and it is important to be exact about which of them is supported by evidence.
      **Time.** Screening, ranking, sizing and writing up the reasons are automated; the adviser reviews a prepared book instead of
      building one. **Discipline.** The same profile always produces the same book on the same day, every exclusion is visible, and
      position sizes follow a stated rule. **Risk communication.** The client is shown a range whose reliability has been measured on
      unseen data. The product does *not* create value by selecting stocks that outperform; Section 11 shows that the evidence does not
      support that claim.""")
    h2("3.2 The quantified case: adviser time saved")
    p(f"""The return-on-investment (ROI) case is built on five inputs, all of them assumptions chosen to describe a mid-sized independent
      practice ([[tab:roi_inputs]]). With {ROI['clients']} clients and {ROI['reviews_per_year']} reviews each a year there are
      {ROI['clients'] * ROI['reviews_per_year']:,} reviews. Saving {saved} minutes on each gives
      {ROI['clients'] * ROI['reviews_per_year']:,} × {saved} ÷ 60 = {BASE['hours']:,.0f} hours a year. Valued at
      Rs {ROI['adviser_rate_inr']:,} an hour this is {inr(BASE['value'])} ({lakh(BASE['value'])}). After an assumed running cost of
      {inr(ROI['annual_cost_inr'])} the net benefit is {inr(BASE['net'])}, and ROI = net benefit ÷ cost = {pct(BASE['roi'], 0)}.""")
    table("roi_inputs", "ROI inputs. All five are assumptions, not measurements", ["Input", "Value", "Basis", "How the team would verify it"], [
        ["Clients per adviser", f"{ROI['clients']}", "Assumed size of a mid-sized independent practice", "Ask pilot advisers for their client count"],
        ["Reviews per client per year", f"{ROI['reviews_per_year']}", "Assumed quarterly review cycle", "Adviser calendar"],
        ["Screening time per review, today", f"{ROI['minutes_before']} min", "Assumed; manual charting and spreadsheet work", "Time-and-motion study on 20 reviews"],
        ["Screening time per review, with the product", f"{ROI['minutes_after']} min", "Assumed; reading a prepared book", "Same study, with the application"],
        ["Value of adviser time", f"Rs {ROI['adviser_rate_inr']:,} / hour", "Assumed billing-equivalent rate", "Practice revenue ÷ adviser hours"],
        ["Annual running cost", inr(ROI["annual_cost_inr"]), "Assumed: hosting, data and maintenance", "Actual hosting and data invoices"],
    ], [4.2, 2.6, 4.9, 4.3])
    h2("3.3 Sensitivity")
    p(f"""Because the inputs are assumed, the conclusion should not depend on any one of them. [[tab:roi_sens]] recomputes ROI for
      smaller and larger practices and for smaller and larger time savings, holding the hourly rate and running cost fixed.
      [[fig:roi]] shows the same calculation as net benefit in rupees. Two points stand out. The case is robust for a practice of
      {ROI['clients']} clients: even if the product saved only 15 minutes per review, less than half the base assumption, ROI would
      still be {pct(roi_case(ROI['clients'], 15)['roi'], 0)}. But the case is weak for a very small practice: with the base saving of
      {saved} minutes, break-even needs about {break_even:.0f} clients, and with 50 clients and a 15-minute saving the product would
      lose money ({pct(roi_case(50, 15)['roi'], 0)}).""")
    sens_clients, sens_minutes = [50, 100, 200, 300, 400], [15, 25, 35, 45]
    table("roi_sens", f"ROI on cost by practice size and minutes saved per review (rate Rs {ROI['adviser_rate_inr']:,}/hour, cost {lakh(ROI['annual_cost_inr'])}/year; base case in bold)",
          ["Clients"] + [f"{m} min saved" for m in sens_minutes] + ["Hours saved (base)"],
          [[str(c)] + [(f"**{pct(roi_case(c, m)['roi'], 0)}**" if (c == ROI["clients"] and m == saved) else pct(roi_case(c, m)["roi"], 0)) for m in sens_minutes]
           + [f"{roi_case(c)['hours']:,.0f}"] for c in sens_clients],
          [2.4, 2.7, 2.7, 2.7, 2.7, 2.8], align="lrrrrr")
    figure("roi", f"""Net annual benefit against number of clients for four levels of time saved per review. Hourly rate
      (Rs {ROI['adviser_rate_inr']:,}) and running cost ({lakh(ROI['annual_cost_inr'])}) are held at the base assumptions.""", 13.5)
    callout("Measured versus assumed", f"""**Measured** (from the hold-out test in this repository): forecast-band coverage of
      {pct(LG['coverage'])} against an {pct(HOLD['target_coverage'], 0)} target; pinball loss, error and rank IC in Section 11; the
      backtest result. **Assumed** (not observed): every input in [[tab:roi_inputs]], and therefore the {BASE['hours']:,.0f} hours,
      the {lakh(BASE['value'])} and the {pct(BASE['roi'], 0)}. **Not claimed:** any improvement in investment returns.""")
    h2("3.4 Market strategy")
    p("""The proposed route to market has three steps. The self-directed investor gets the research tool free, which builds an audience
      and a body of user feedback. The adviser pays for a seat, because the adviser is the user for whom saved hours convert directly
      into money and who needs client-ready evidence sheets (the CSV export and the evidence ledger). Brokers are a distribution channel
      and not a customer in the first phase: the application already contains adapters for the Upstox and Zerodha market-data feeds, so
      it can sit beside an existing brokerage login. This strategy is a proposal. No pricing test, customer interview programme or
      pilot has been run, and none is claimed.""")

    # ------------------------------------------------------- 4 architecture
    h1("4. System architecture and repository")
    p("""The system has two halves that meet at a small set of files ([[fig:architecture]]). The **training pipeline** runs offline, on
      demand, and takes about a minute. The **product** reads what the pipeline wrote and never trains anything. Keeping them apart
      means the web application cannot change the model by accident, and the numbers a user sees are exactly the numbers that were
      validated.""")
    figure("architecture", """System architecture. Training (centre) turns the committed data snapshot into two artefacts: a report of
      every metric and a bundle of fitted models. The product (right) reads the bundle; this report is generated from the metrics file.""")
    p("""`src/data.py` holds everything about data: the universe, the download and snapshot functions, the technical indicators,
      fractional differencing, the feature-selection routine and the time-series split functions. `src/model.py` holds everything about
      learning: the loss function and metrics, the quantile models, the random search, the baselines, the hold-out evaluation, the
      backtest, SHAP and the error analysis; its `train()` function writes the artefacts. `src/portfolio_advisor/core.py` is the
      decision engine that converts forecasts into ratings and position sizes (Section 14). `app.py` is the Streamlit interface and
      `api_app.py` exposes the same engine as a JSON web service for other software.""")
    table("layout", "Repository layout", ["Path", "Role"], [
        ["`src/data.py`", "Universe, price download and snapshot, 19 candidate features, fractional differencing, feature selection, purged time-series splits"],
        ["`src/model.py`", "Pinball loss and metrics, random search, walk-forward cross-validation, baselines, final quantile models, hold-out test, backtest, SHAP, error analysis, export"],
        ["`app.py`", "Streamlit product: profile intake, seven tabs, analyst chat"],
        ["`requirements.txt`, `pyproject.toml`", "Pinned dependency ranges; Python 3.11 to 3.13"],
        ["`src/portfolio_advisor/core.py`", "Decision engine: pillar scores, profile gates, ratings, HRP sizing, whole-share allocation"],
        ["`src/portfolio_advisor/live.py`, `chat.py`", "Optional live-quote adapters (Upstox, Zerodha, Yahoo) and the grounded chat client"],
        ["`api_app.py`", "Optional FastAPI service exposing the same engine as JSON endpoints"],
        ["`data/snapshot/`", "Committed prices and fundamentals, so that training is reproducible without a network connection"],
        ["`reports/metrics.json`, `reports/cv_results.csv`", "Every number in this report; one row per trial and fold"],
        ["`tests/`", "22 automated tests, including the leakage checks"],
        ["`scripts/build_report.py`", "Generates this report and its figures from the two files in `reports/`"],
        ["`docs/`", "This report, the slide deck, and the vibe coding log"],
        ["`Dockerfile`, `docker-compose.yml`, `.github/workflows/ci.yml`", "Container images, two-service deployment, and continuous integration"],
    ], [5.6, 10.4])

    # --------------------------------------------------------------- 5 data
    h1("5. Data")
    h2("5.1 Universe, period and source")
    p(f"""The universe is {D['tickers']} NSE-listed large caps, five from each of {D['sectors']} sectors ([[tab:universe]]). Fixing five
      per sector stops the model from learning mainly about banks and information-technology companies, which dominate Indian indices
      by weight. Prices are daily open, high, low, close and volume from Yahoo Finance through the `yfinance` library, adjusted for
      splits and dividends, downloaded from {'2017-01-01'}. After the indicators have enough history to be computed, the usable panel
      runs from {D['first_date']} to {D['last_date']} and contains {D['panel_rows']:,} rows, where one row is one stock on one trading
      day. The downloaded data are committed to the repository as a compressed snapshot, so anyone who clones the repository trains on
      exactly the same data.""")
    table("universe", "The 60-stock universe by sector (NSE tickers)", ["Sector", "Tickers"], universe_rows(), [5.2, 10.8])
    h2("5.2 The label")
    p(f"""The quantity to be predicted, the label, is the **forward return**: the percentage change in the adjusted close from day *t*
      to day *t* + {CFG['HORIZON']} trading days. It is a continuous number, so the task is regression, and it is defined for every row
      except the most recent {CFG['HORIZON']} days, whose future is not yet known. Those {D['holdout_rows_latest_unlabelled']:,} newest
      rows are kept so that the application can score today's market, and they are never used for training or for any metric.""")
    h2("5.3 Cleaning and preprocessing")
    p("""Four steps prepare the data. (1) Tickers with fewer than 250 days of prices are dropped at download. (2) Infinite values produced
      by divisions (for example a zero 14-day range) are converted to missing. (3) Rows missing any of the 19 candidate features are
      dropped; these are almost entirely the warm-up period at the start of each stock's history, when a 252-day high or a fractional
      difference cannot yet be computed. No values are imputed, because inventing a price indicator would add noise where we can simply
      wait for data. (4) Every feature is a return, a ratio or a bounded oscillator, so a stock trading at Rs 300 and one at Rs 30,000
      are directly comparable without any rescaling of price levels.""")
    h2("5.4 How the rows are divided")
    p(f"""The dates are divided once, by time, before any modelling ([[tab:rows]]). The earliest 80% of dates are the **development
      set**, used for feature selection, tuning and cross-validation. The latest 20% are the **hold-out set**, used once for the final
      test. The {GAP} trading days immediately before the hold-out are discarded so that no development label can reach into the
      hold-out period (Section 9 explains why).""")
    table("rows", "How the panel is divided", ["Part", "Dates", "Rows", "Used for"], [
        ["Development", f"{D['first_date']} to {D['development_end']}", f"{D['development_rows']:,}", "Feature selection, tuning, cross-validation, final fit"],
        ["Purged gap", f"{GAP} trading days", f"{purge_rows:,}", "Discarded"],
        ["Hold-out, scored", f"{D['holdout_start']} to {D['holdout_end_scored']}", f"{D['holdout_rows_scored']:,}", "Final test, SHAP, error analysis, backtest"],
        ["Latest, no label yet", f"last {CFG['HORIZON']} trading days to {D['last_date']}", f"{D['holdout_rows_latest_unlabelled']:,}", "Live forecasts in the application only"],
        ["Total", "", f"{D['panel_rows']:,}", ""],
    ], [3.4, 5.0, 2.2, 5.4], align="llrl")
    h2("5.5 Survivorship bias")
    p("""The universe was chosen in 2026 from companies that are large today. Companies that were large in 2018 and later shrank, merged
      or were delisted are absent, and companies that grew into large caps are included for their whole history. This is survivorship
      bias. It flatters any historical return figure for the universe as a whole, which is one reason the backtest benchmark is the
      same 60 stocks held equally and not an outside index: both the strategy and the benchmark carry the same bias, so their difference
      is a fairer measure than either level.""")

    # ----------------------------------------------------------- 6 features
    h1("6. Feature engineering")
    p(f"""A feature is an input column the model can use. All {n_candidates} candidates are computed only from prices and volumes up to
      and including day *t*, so each is known at the moment a forecast would be made. Sixteen describe the individual stock and three
      describe the market as a whole on that day ([[tab:features]]). The market features are simple cross-sectional averages over the
      {D['tickers']} stocks; they let the model behave differently in calm and turbulent markets.""")
    selected_set = set(M["feature_selection"]["selected"])
    table("features", f"The {n_candidates} candidate features", ["Feature", "How it is computed", "What it measures", "Kept"],
          [[f"`{name}`", how, what, "Yes" if name in selected_set else "No"] for name, how, what in FEATURES], [2.7, 6.9, 5.3, 1.1], size=8)
    h2("6.1 Why prices cannot be used directly: stationarity")
    p("""A share price trends. Its average level in 2019 is different from its average level in 2025, so a rule learned at one price
      level does not transfer to another. Statisticians call a series **stationary** when its mean and variance do not drift over time,
      and most learning algorithms implicitly assume their inputs are stationary. The usual cure is to take the day-to-day change
      (first difference, or return). That makes the series stationary but throws away almost all of its memory: a return says nothing
      about whether the stock is high or low relative to its own past.""")
    h2("6.2 Fractional differencing from first principles")
    p("""Ordinary differencing computes x(t) − x(t−1): a weight of 1 on today and −1 on yesterday. Fractional differencing generalises
      this to an order *d* between 0 and 1. The weights come from the binomial series: the first is 1, and each later weight is the
      previous one multiplied by −(d − k + 1) ÷ k. For d = 1 the sequence is 1, −1, 0, 0, … which is the ordinary difference; a unit
      test in the repository confirms this. For d = 0.4 the sequence is 1, −0.40, −0.12, −0.064, … and decays slowly, so the result
      still depends on prices from many weeks back. The smaller *d* is, the more memory survives; the larger it is, the more of the
      trend is removed. The code stops adding weights once they fall below 0.0001 in absolute size.""")
    p("""The right *d* is the smallest one that makes the series stationary. To decide, the pipeline uses the **Augmented Dickey-Fuller
      (ADF) test**. Its null hypothesis is that the series has a unit root, meaning that shocks never fade and the series wanders
      without a fixed mean. A p-value below 0.05 rejects that hypothesis. For each stock the pipeline tries d = 0, 0.1, 0.2, … 1.0 on
      the log price and keeps the first value that passes; the resulting series is the feature `fracdiff_close`. Crucially, *d* is
      chosen using development-period prices only (the `fit_until` argument in `stock_features`), so the hold-out period has no
      influence on how the feature is built. This technique follows López de Prado (2018).""")

    # ---------------------------------------------------- 7 feature selection
    h1("7. Feature selection")
    h2("7.1 Why select at all")
    p("""Many technical indicators are different arithmetic on the same information. The clearest case in our list is `Williams_R`, which
      is exactly `Stoch_K` minus 100. Feeding near-duplicates to a model adds no information, makes importance scores harder to read
      because credit is split between twins, and gives a flexible model more ways to fit noise. Selection therefore aims for a small
      set in which each feature says something the others do not.""")
    h2("7.2 The method: relevance minus redundancy")
    p("""**Mutual information** between two variables measures how much knowing one reduces uncertainty about the other. It is zero
      when they are independent and, unlike correlation, it also detects non-linear relationships (for example, "very high and very low
      volatility both precede large moves"). The pipeline uses scikit-learn's nearest-neighbour estimator, which reports the value in
      natural units (nats).""")
    p(f"""Selection is greedy. Step 1 picks the feature with the highest mutual information with the label (its *relevance*). Every later
      step scores each remaining feature as relevance minus 0.5 times its *redundancy*, where redundancy is its average mutual
      information with the features already chosen, and adds the best one. This repeats until {n_selected} features are chosen. The idea
      is that of the minimum-redundancy maximum-relevance method of Peng, Long and Ding (2005). The calculation runs on a random sample
      of 20,000 development rows; the hold-out period is never seen.""")
    figure("selection", f"""Feature selection. Bars show each feature's mutual information with the 21-day forward return; dots show the
      redundancy of each selected feature with those chosen before it (the first pick, `mkt_vol_21`, has none). Several high-relevance
      features were rejected because they duplicate one already chosen.""", 13.0)
    table("selection", "Feature-selection table, in the order features were chosen (mutual information in nats)",
          ["Order", "Feature", "Relevance", "Redundancy", "Score", "Selected"],
          [[str(i + 1) if r["selected"] else "", f"`{r['feature']}`", f"{r['relevance']:.5f}",
            "" if r["redundancy"] != r["redundancy"] else f"{r['redundancy']:.5f}", "" if r["score"] != r["score"] else f"{r['score']:+.5f}",
            "Yes" if r["selected"] else "No"] for i, r in enumerate(sel_table)], [1.6, 4.0, 2.7, 2.7, 2.7, 2.3], align="llrrrl")
    h2("7.3 Reading the result")
    rel = {r["feature"]: r for r in sel_table}
    p(f"""Three observations follow from [[tab:selection]] and [[fig:selection]]. First, every relevance value is small: the largest is
      {rel['mkt_vol_21']['relevance']:.3f} nats. No feature knows much about next month's return, which is the first honest signal that
      this is a hard prediction problem. Second, the penalty does what it was designed to do. `ATR_14`, `vol_21` and `mkt_ret_21` have
      higher relevance than several selected features but were left out because volatility and market information were already
      represented by `mkt_vol_21`, `vol_63` and `rel_ret_21`; of the Stoch_K and Williams_R twins only one was kept. Third, the method
      has a weakness that we report: `OBV_delta` has an estimated relevance of {rel['OBV_delta']['relevance']:.5f} and was still
      chosen ninth, because once the informative features are taken, a feature that is merely *different* scores better than one that
      is informative but redundant. The SHAP analysis in Section 12 confirms that it contributes almost nothing to forecasts. A stricter
      rule (stop when the score turns negative) would have kept only two features.""")
    h2("7.4 Why company fundamentals were removed from the model")
    p("""The original notebook also fed the model seven fundamental ratios: price-to-earnings, price-to-book, return on equity,
      debt-to-equity, net profit margin, dividend yield and revenue growth. Yahoo Finance supplies only **today's** value of each. The
      notebook copied today's value onto every historical row, so a training row dated 2019 "knew" the company's 2026 profitability.
      That is **look-ahead leakage**: information from the future reaching the model. It inflates every test score, because a company
      that is highly profitable in 2026 is, on average, one whose share price rose between 2019 and 2026.""")
    p("""The fundamentals were therefore removed from the list of candidate features, and a unit test
      (`test_fundamentals_are_not_model_inputs`) fails the build if any feature whose name begins with `f_` is added back. They are
      still used, legitimately, in one place: the application's separate **fundamental score** for *today's* decision (Section 14),
      where using today's ratios involves no look-ahead. The application and the artefact manifest both state that this score is not
      part of the validated forecast.""")

    # ------------------------------------------------------------ 8 modelling
    h1("8. Modelling")
    h2("8.1 Quantile regression and pinball loss")
    p("""Ordinary regression minimises squared error, and the forecast that minimises squared error is the average outcome. An adviser
      needs more than the average: the client's question is "how bad could this be?". **Quantile regression** (Koenker and Bassett,
      1978) forecasts a chosen percentile of the outcome instead. A P10 forecast is the level the return should fall below only 10% of
      the time; P50 is the median; P90 is the level it should exceed only 10% of the time. The interval from P10 to P90 should
      therefore contain 80% of outcomes.""")
    p("A model is trained to forecast quantile *q* by minimising the **pinball loss**. For an actual return *y* and a forecast *ŷ*:")
    eq("L(y, ŷ) = q × (y − ŷ)   if y ≥ ŷ          L(y, ŷ) = (1 − q) × (ŷ − y)   if y < ŷ")
    p("""It is an absolute error with unequal prices on the two sides ([[fig:pinball]]). For q = 0.1, an outcome *above* the forecast
      costs only 0.1 per unit, while an outcome *below* it costs 0.9 per unit. The cheapest place to put the forecast is therefore low,
      at the point where only 10% of outcomes fall beneath it. For q = 0.5 both sides cost the same and the loss is half the absolute
      error, whose minimiser is the median.""")
    figure("pinball", """Pinball loss for the three quantiles as a function of the forecast error. The asymmetric slopes are what push
      each model towards its own percentile.""", 11.5)
    p("""**A worked example.** Suppose the three forecasts for a stock are P10 = −6%, P50 = +1% and P90 = +9%, and the return turns out
      to be +2%. For P10 the outcome is 8 points above the forecast, so the loss is 0.1 × 0.08 = 0.008. For P50 it is 1 point above:
      0.5 × 0.01 = 0.005. For P90 it is 7 points below: (1 − 0.9) × 0.07 = 0.007. The mean over the three quantiles is 0.0067. If the
      return had instead been −10%, four points *below* the P10 forecast, the P10 loss alone would be 0.9 × 0.04 = 0.036: a breach of
      the pessimistic case is punished heavily. The mean pinball loss over the three quantiles is the single number used to tune and
      compare models in this report; lower is better, and it is in the same units as the return.""")
    h2("8.2 Gradient-boosted trees and LightGBM in plain terms")
    p("""A **decision tree** asks a sequence of yes/no questions about the features ("is market volatility above 1.4%?") and gives every
      row that reaches the same end-point, or leaf, the same prediction. One small tree is a crude model. **Gradient boosting**
      (Friedman, 2001) builds a strong model from many crude ones: start from a constant prediction, look at where the current model is
      wrong, fit a small tree to those errors, add a fraction of that tree's output to the model, and repeat. The fraction is the
      *learning rate*. With a quantile objective, "where the model is wrong" is defined by the pinball loss, and the starting constant
      is the corresponding percentile of the training labels.""")
    p("""**LightGBM** (Ke et al., 2017) is an efficient implementation of this idea. Two design choices matter here. It sorts each
      feature's values into a fixed number of bins before training, which makes finding split points fast. And it grows trees *leaf-wise*:
      it always splits the leaf whose split reduces the loss most, so tree size is controlled by the number of leaves (`num_leaves`)
      and not by depth. We chose gradient-boosted trees because they handle non-linear effects and interactions between features without
      manual specification, need no feature scaling, are robust to outliers in the inputs, support the quantile objective directly, and
      can be explained exactly with SHAP (Section 12).""")
    h2("8.3 Why three models, and the quantile-crossing fix")
    p("""Each quantile has its own loss function, so each needs its own model: three separate LightGBM regressors with the objective set
      to `quantile` and `alpha` set to 0.1, 0.5 and 0.9. They share the same features and hyperparameters. Because they are trained
      independently, nothing forces their outputs to be ordered; occasionally the P10 model can predict a higher value than the P50
      model for the same row, which is logically impossible. This is known as **quantile crossing**. The function `predict_quantiles`
      fixes it by sorting the three predictions for every row, so P10 ≤ P50 ≤ P90 always holds; a unit test feeds it deliberately
      crossed values and checks the ordering.""")
    h2("8.4 Baselines")
    p(f"""A model's score means nothing without something to compare it with. Two baselines are run on exactly the same folds. The
      **naive baseline** ignores every feature: for each quantile it predicts the corresponding percentile of all returns in the
      training period, the same three numbers for every stock on every day. It answers the question "do the features add anything?".
      The **Ridge baseline** is a linear regression with an L2 penalty (Hoerl and Kennard, 1970) that predicts the expected return.
      It answers "does a non-linear model add anything over a straight-line one?". Its penalty strength *alpha* was chosen from six
      values between 0.1 and 10,000 by cross-validated mean absolute error; the largest, {RIDGE['settings'].split('=')[1]}, won, which
      itself says the data support only heavily shrunk coefficients. Ridge gives a single forecast and no quantiles, so it is compared
      on error and rank IC but has no pinball loss or coverage.""")
    h2("8.5 Scaling: needed for Ridge, not for trees")
    p("""Ridge penalises the size of its coefficients, and the size of a coefficient depends on the units of its feature: RSI runs from
      0 to 100 while a daily volatility is around 0.02. Without rescaling, the penalty would fall almost entirely on the small-unit
      features. Ridge is therefore run inside a scikit-learn pipeline with a `StandardScaler`, which subtracts each feature's mean and
      divides by its standard deviation. The scaler is **refitted on the training rows of every fold**; fitting it once on all the data
      would let validation-period statistics leak into training.""")
    p("""Trees need no scaling. A split asks only whether a feature is above or below a threshold, so any order-preserving transformation
      of the feature (rescaling, taking logs) produces exactly the same partitions of the rows. LightGBM therefore receives the raw
      features. This is a deliberate decision and not an omission.""")

    # ----------------------------------------------------------- 9 validation
    h1("9. Validation design")
    h2("9.1 Why ordinary cross-validation is wrong for this data")
    p("""Standard k-fold cross-validation shuffles the rows and holds out a random fifth at a time. With time-ordered market data that
      is wrong twice over. It trains on the future to predict the past, which no real user can do. And neighbouring rows are nearly
      duplicates: the 21-day return starting on Monday and the one starting on Tuesday share 20 of their 21 days, so a model can score
      well on a shuffled test simply by having seen an almost identical row in training.""")
    h2("9.2 Purged walk-forward cross-validation")
    p(f"""The pipeline uses **walk-forward** validation with an expanding window. The development dates are cut into six equal blocks.
      Fold 1 trains on block 1 and validates on block 2; fold 2 trains on blocks 1 and 2 and validates on block 3; and so on, giving
      {CFG['N_SPLITS']} folds ([[fig:folds]], [[tab:folds]]). Training data always precede validation data, exactly as in live use, and
      each fold tests the model on a different market period, including the sharp fall and recovery of early 2020 in fold 1.""")
    figure("folds", f"""Walk-forward folds and the final hold-out. The thin marigold strip before every validation block is the purged
      gap of {GAP} trading days.""", 13.5)
    table("folds", "Cross-validation folds", ["Fold", "Training rows", "Validation rows", "Validation from", "Validation to"],
          [[str(f["fold"]), f"{f['train_rows']:,}", f"{f['val_rows']:,}", f["val_start"], f["val_end"]] for f in FOLDS]
          + [["Final", f"{D['development_rows']:,}", f"{D['holdout_rows_scored']:,}", D["holdout_start"], D["holdout_end_scored"]]],
          [2.0, 3.4, 3.4, 3.6, 3.6], align="lrrll")
    h2(f"9.3 Why the gap is {GAP} trading days")
    p(f"""A training row dated day *t* carries a label that depends on the price at *t* + {CFG['HORIZON']}. If the last training day
      were the day before validation starts, its label would be computed from prices inside the validation period: the model would be
      trained on information from the very days it is about to be tested on. **Purging** removes those rows: the last
      {CFG['HORIZON']} trading days before each validation block are dropped from training, so the latest training label ends before
      validation begins. A further **embargo** of {CFG['EMBARGO']} trading days is added as a safety margin, because volatility and
      market mood persist for days and a label ending the day before validation still describes almost the same market.
      {CFG['HORIZON']} + {CFG['EMBARGO']} = {GAP} trading days, about six calendar weeks. Both ideas are from López de Prado (2018).
      The gap costs data, which is why fold 1 trains on only {FOLDS[0]['train_rows']:,} rows.""")
    h2("9.4 The hold-out period")
    p(f"""Cross-validation is used many times: to compare {M['validation']['lightgbm_trials']} hyperparameter settings and six Ridge
      penalties. Any number that has been used to choose among options is optimistic, because the winner is partly the luckiest. The
      hold-out period, {D['holdout_start']} to {D['holdout_end_scored']} ({D['holdout_rows_scored']:,} rows over {LG['ic_days']} trading
      days), played no part in any choice. After tuning, the three final models were refitted once on all development data with the
      chosen settings and scored once on the hold-out. Those are the numbers a user should expect in live use, and they are the ones
      quoted in the executive summary.""")
    h2("9.5 Leakage controls")
    p("[[tab:leakage]] lists every route by which future information could reach the model and the control that closes it.")
    table("leakage", "Leakage controls", ["Risk", "Control", "Where enforced"], [
        ["Training labels overlap the validation period", f"{GAP}-day purge and embargo before every validation block", "`walk_forward_splits`; test `test_walk_forward_folds_never_train_on_the_future`"],
        ["Development labels overlap the hold-out period", f"The same {GAP}-day gap before the hold-out", "`holdout_dates`; test `test_holdout_is_purged_from_development`"],
        ["Training on the future", "Expanding window: validation always follows training", "`walk_forward_splits`"],
        ["Today's fundamentals copied onto past dates", "Fundamentals removed from model inputs", "`CANDIDATE_FEATURES`; test `test_fundamentals_are_not_model_inputs`"],
        ["Feature selection sees the hold-out", "Mutual information computed on development rows only", "`train()` in `src/model.py`"],
        ["Differencing order tuned on the hold-out", "*d* chosen from prices up to the development end date", "`fit_until` in `stock_features`"],
        ["Scaler fitted on validation data", "`StandardScaler` inside the pipeline, refitted per fold", "`cross_validate_baselines`"],
        ["Hyperparameters tuned on the hold-out", "Search scored only on development folds", "`tune_lightgbm`"],
        ["Features use future prices", "Every indicator uses rolling windows that end at day *t*", "`stock_features`"],
    ], [5.0, 5.4, 5.6], size=8)
    p("""One residual imperfection should be stated. Feature selection and the differencing order are fitted once on the whole
      development set, and the cross-validation folds are then drawn inside that set. The fold scores are therefore very slightly
      optimistic, because the selection step has seen the fold's validation rows. The hold-out is not affected, and it is the hold-out
      on which our conclusions rest.""")

    # ------------------------------------------------------------------ 10 HPO
    h1("10. Hyperparameter optimisation")
    h2("10.1 Search space and method")
    p(f"""Hyperparameters are the settings fixed before training. Seven were searched ([[tab:space]]). Their grid contains
      {int(np.prod([len(v) for v in M['validation']['search_space'].values()])):,} combinations, too many to try exhaustively, so the
      pipeline uses **random search**: it draws {M['validation']['lightgbm_trials'] - 1} combinations at random (with a fixed seed of
      {CFG['SEED']}, so the draw is repeatable) and adds the original notebook settings as trial 0, for
      {M['validation']['lightgbm_trials']} settings in all. Random search covers a large space more efficiently than a grid when only a
      few of the parameters matter (Bergstra and Bengio, 2012). Each setting is trained on all {CFG['N_SPLITS']} folds for all three
      quantiles: {M['validation']['lightgbm_trials']} × {CFG['N_SPLITS']} × 3 = {M['validation']['lightgbm_fits']} model fits, run in
      parallel. The score for a setting is its mean pinball loss across the three quantiles, averaged over the five folds.""")
    table("space", "Hyperparameter search space", ["Parameter", "Values tried", "What it controls", "Notebook", "Chosen"],
          [[f"`{k}`", ", ".join(f"{v:g}" for v in values), PARAM_NOTES[k], f"{SEL['notebook_params'][k]:g}", f"**{params[k]:g}**"]
           for k, values in M["validation"]["search_space"].items()], [3.2, 3.0, 6.4, 1.7, 1.7], align="lllrr", size=8)
    h2("10.2 What the search found")
    top, bottom = TRIALS.head(5), TRIALS.tail(5)

    def trial_row(rank: int, r) -> list[str]:
        tag = " (chosen)" if int(r["trial"]) == BEST_TRIAL else " (notebook)" if int(r["trial"]) == 0 else ""
        return [str(rank), f"{int(r['trial'])}{tag}", f"{r['mean']:.5f}", f"{r['std']:.5f}", f"{int(r['n_estimators'])}", f"{r['learning_rate']:g}",
                f"{int(r['num_leaves'])}", f"{int(r['min_child_samples'])}", f"{r['subsample']:g}", f"{r['colsample_bytree']:g}", f"{r['reg_lambda']:g}"]

    trial_header = ["#", "Trial", "Pinball mean", "Fold s.d.", "Trees", "Learn. rate", "Leaves", "Min. leaf rows", "Row share", "Feature share", "L2"]
    trial_widths = [0.8, 2.5, 1.9, 1.6, 1.2, 1.4, 1.5, 1.6, 1.2, 1.4, 0.9]
    table("trials_short", "Best five and worst five settings by cross-validated pinball loss (full list in Appendix B)", trial_header,
          [trial_row(i + 1, r) for i, (_, r) in enumerate(top.iterrows())] + [["…"] + [""] * 10]
          + [trial_row(len(TRIALS) - 4 + i, r) for i, (_, r) in enumerate(bottom.iterrows())], trial_widths, align="llrrrrrrrrr", size=8)
    figure("trials", """All 25 settings ranked by mean pinball loss. Each dot is the mean over five folds and each vertical line runs
      from the best to the worst fold. The chosen setting is indigo, the original notebook setting marigold, and the horizontal line is
      the naive baseline.""", 14.0)
    p(f"""The chosen setting is trial {BEST_TRIAL}: {params_text(params)}. Its cross-validated pinball loss is
      {f5(TUNED['pinball_mean'])} (standard deviation across folds {f5(TUNED['pinball_std'])}), against {f5(NOTEBOOK['pinball_mean'])}
      for the notebook settings: an improvement of {pct(SEL['pinball_improvement_vs_notebook'])}. The notebook setting ranks
      {int(TRIALS.index[TRIALS['trial'] == 0][0]) + 1} of {len(TRIALS)}.""")
    h2("10.3 Why small trees won")
    small = TRIALS.head(8)
    p(f"""The pattern in [[tab:trials_short]] and Appendix B is consistent. All eight best settings use
      {int(small['n_estimators'].min())} or {int(small['n_estimators'].max())} trees, and each is constrained in at least one way:
      very few leaves (4 or 8), or a very large minimum leaf size (1,000 rows), or a low learning rate. The worst settings have many trees (400 or 800), a high
      learning rate, large trees, or several of these together; no setting with 400 or more trees ranks in the top twelve. In other words, the more freedom the model was given, the worse it
      did out of sample. This is the signature of a **low signal-to-noise** problem: a 31-leaf tree has enough flexibility to memorise
      accidental patterns in past returns, and those patterns do not repeat.""")
    p(f"""The chosen model is small in a precise sense. Each tree has {params['num_leaves']:g} leaves, so it asks at most three
      questions and can combine at most three features. With {params['n_estimators']:g} trees and a learning rate of
      {params['learning_rate']:g}, the model starts from the historical percentile and adds one hundred corrections, each scaled down to
      2% of what the tree proposed. It is therefore a cautious adjustment of the naive forecast, not a replacement for it, which is
      exactly why its cross-validated score lands so close to the naive baseline (Section 11). `colsample_bytree` =
      {params['colsample_bytree']:g} offers each tree a random half of the ten features, and `subsample` = {params['subsample']:g} a
      random 80% of rows, which makes the trees differ from one another and averages out noise.""")
    p(f"""Two cautions apply. The top five settings are separated by {TRIALS.iloc[4]['mean'] - TRIALS.iloc[0]['mean']:.5f} in pinball
      loss, while the standard deviation across folds is about {TUNED['pinball_std']:.3f}: the ranking among them is not statistically
      meaningful, and we chose the lowest mean because a rule fixed in advance is better than a choice made after looking. And
      [[fig:trials]] shows that which *period* is being predicted moves the loss far more than which *setting* is used.""")

    # -------------------------------------------------------------- 11 results
    h1("11. Results")
    h2("11.1 Cross-validation: four candidates on the same folds")

    def opt(value, fmt="{:.5f}") -> str:
        return "n/a" if value is None else fmt.format(value)

    table("cv", "Cross-validation results (mean over five folds; ± is the standard deviation across folds)",
          ["Model", "Pinball loss", "MAE", "Rank IC", "Coverage", "Direction right"],
          [[("**" + c["model"] + "**") if c is TUNED else c["model"],
            "n/a" if c["pinball_mean"] is None else f"{c['pinball_mean']:.5f} ± {c['pinball_std']:.5f}",
            f"{c['mae_mean']:.4f}", f"{c['ic_mean']:+.3f} ± {c['ic_std']:.3f}",
            "n/a" if c["coverage_mean"] is None else pct(c["coverage_mean"]), pct(c["directional_accuracy_mean"])] for c in M["candidates"]],
          [4.6, 3.3, 1.5, 2.7, 1.9, 2.0], align="lrrrrr", size=8)
    figure("cv", """Pinball loss in each validation fold for the three quantile candidates. Ridge is omitted because it produces no
      quantiles.""", 13.5)
    p(f"""[[tab:cv]] and [[fig:cv]] give three findings. **Tuning mattered.** The notebook settings were not merely sub-optimal; they
      were worse than using no features at all ({f5(NOTEBOOK['pinball_mean'])} against {f5(NAIVE['pinball_mean'])}), and their P10 to
      P90 band contained only {pct(NOTEBOOK['coverage_mean'])} of outcomes instead of 80%: an over-fitted model is over-confident.
      **The tuned model ties the naive baseline.** Its loss of {f5(TUNED['pinball_mean'])} is
      {pct(abs(SEL['pinball_improvement_vs_naive']))} *worse* than the naive {f5(NAIVE['pinball_mean'])}, a gap of
      {TUNED['pinball_mean'] - NAIVE['pinball_mean']:.5f} that is roughly one-fiftieth of the fold-to-fold standard deviation. The
      correct reading is "no detectable difference". **Ranking skill is not established in cross-validation.** The tuned model's mean
      rank IC is {sic(TUNED['ic_mean'])} with a standard deviation of {TUNED['ic_std']:.3f}: positive in two folds and negative in
      three. The naive baseline has the highest directional accuracy ({pct(NAIVE['directional_accuracy_mean'])}) for a simple reason:
      its median forecast is always slightly positive, and in a rising market "up" is the right call more often than not.""")
    h2("11.2 Hold-out test")
    table("holdout", f"Hold-out results, {D['holdout_start']} to {D['holdout_end_scored']} ({D['holdout_rows_scored']:,} rows)",
          ["Measure", "LightGBM (tuned)", "Naive", "Ridge", "Better is"], [
              ["Pinball loss (mean of 3 quantiles)", f"**{f5(LG['pinball'])}**", f5(NV["pinball"]), "n/a", "lower"],
              ["P10 to P90 coverage (target 80%)", f"**{pct(LG['coverage'])}**", pct(NV["coverage"]), "n/a", "closer to 80%"],
              ["Average band width", f"**{LG['band_width'] * 100:.1f} pp**", f"{NV['band_width'] * 100:.1f} pp", "n/a", "narrower, if coverage holds"],
              ["MAE of the median forecast", f4(LG["mae"]), f4(NV["mae"]), f"**{f4(RG['mae'])}**", "lower"],
              ["RMSE of the median forecast", f"**{f4(LG['rmse'])}**", f4(NV["rmse"]), f4(RG["rmse"]), "lower"],
              ["Direction right", pct(LG["directional_accuracy"]), pct(NV["directional_accuracy"]), f"**{pct(RG['directional_accuracy'])}**", "higher"],
              ["Mean daily rank IC", sic(LG["ic_mean"]), "0 (no ranking)", f"**{sic(RG['ic_mean'])}**", "higher"],
              ["Days with positive rank IC", f"{pct(LG['ic_positive_share'])} of {LG['ic_days']}", "n/a", f"**{pct(RG['ic_positive_share'])}** of {RG['ic_days']}", "higher"],
          ], [5.2, 3.0, 2.4, 2.6, 2.8], align="lrrrl", size=8)
    p(f"""On data the model never saw, the picture is slightly more favourable but tells the same story ([[tab:holdout]]). The tuned
      model's pinball loss is {pct(hold_gain)} lower than the naive baseline's. More usefully, it achieves this with a band that is
      {pct(band_cut)} narrower ({LG['band_width'] * 100:.1f} against {NV['band_width'] * 100:.1f} percentage points) while its coverage
      is closer to the 80% target. A narrower band that is still honest is a real, if modest, gain for the user. On point accuracy the
      three models are indistinguishable: mean absolute errors of {f4(LG['mae'])}, {f4(NV['mae'])} and {f4(RG['mae'])} differ in the
      fourth decimal place, and the direction of the move is called correctly {pct(LG['directional_accuracy'])} of the time, barely
      better than a coin. The simple linear Ridge model is marginally ahead of LightGBM on mean absolute error and rank IC, which is further evidence
      that the data do not reward model complexity.""")
    h2("11.3 Coverage and calibration")
    p(f"""A forecast band is **calibrated** when it contains the share of outcomes it promises. The P10 to P90 band promises 80%. On the
      hold-out it contained {pct(LG['coverage'])} ([[fig:coverage]]): slightly conservative, erring on the side of a band that is a
      little too wide, which is the safer direction when the band is used to warn a client about downside. Across the
      cross-validation folds coverage ranged from {pct(_lgbm_fold(BEST_TRIAL, 'coverage').min())} to
      {pct(_lgbm_fold(BEST_TRIAL, 'coverage').max())}. Calibration therefore holds on average but not in every market period: a band
      learned from a calm history is too narrow when a turbulent period follows (fold 1), and too wide in the reverse case (fold 4).
      This is the product's strongest measured property, and the claim attached to it is deliberately limited: "about four outcomes in
      five have fallen inside the range", not "the range is guaranteed".""")
    figure("coverage", """Share of real outcomes inside the P10 to P90 band, by cross-validation fold and on the hold-out. The notebook
      settings were systematically over-confident; the tuned model tracks the 80% target on average.""", 13.5)
    h2("11.4 Rank information coefficient")
    p(f"""The **rank information coefficient** (rank IC) asks a different question from error: on a given day, did the stocks the model
      ranked highest go on to do best? It is the Spearman rank correlation, across the {D['tickers']} stocks, between the day's median
      forecasts and the returns that followed, computed for each day and then averaged. Zero means no ordering skill and +1 a perfect
      ordering. In equity forecasting, values of a few hundredths are typical of real signals, so small numbers are expected; what
      matters is whether they are reliably positive.""")
    p(f"""On the hold-out the mean rank IC is {sic(LG['ic_mean'])}, positive on {pct(LG['ic_positive_share'])} of {LG['ic_days']} days.
      That is encouraging, but two facts temper it. The same statistic was {sic(TUNED['ic_mean'])} ± {TUNED['ic_std']:.3f} across the
      cross-validation folds, so a single positive period cannot be told apart from chance. And daily values on overlapping 21-day
      returns are strongly correlated with each other, so {LG['ic_days']} days are far fewer than {LG['ic_days']} independent
      observations. We therefore describe ranking skill as *unproven*, not as weak-but-present.""")
    h2("11.5 Do the ratings mean anything? Realised returns by signal")
    p(f"""The model also emits a five-level signal from the day's ranking. [[tab:signals]] shows what followed each signal on the
      hold-out. The three middle levels are ordered as intended: BUY rows returned {spct(signals['BUY']['mean_forward_return'], 2)} on
      average, HOLD {spct(signals['HOLD']['mean_forward_return'], 2)} and SELL {spct(signals['SELL']['mean_forward_return'], 2)}. The
      two extreme levels are not: the {signals['STRONG_BUY']['rows']:,} STRONG_BUY rows *lost*
      {pct(abs(signals['STRONG_BUY']['mean_forward_return']), 2)} on average and STRONG_SELL rows gained. The "strong" labels require
      both an extreme forecast and a narrow band, a combination that proved unreliable. This is why the application does not pass the
      raw model signal to the user: the decision engine in Section 14 requires agreement from other evidence before issuing a
      STRONG BUY.""")
    table("signals", "Hold-out outcomes by raw model signal", ["Model signal", "Rows", "Mean 21-day return", "Share with a positive return"],
          [[s, f"{signals[s]['rows']:,}", spct(signals[s]["mean_forward_return"], 2), pct(signals[s]["share_positive"])]
           for s in ["STRONG_BUY", "BUY", "HOLD", "SELL", "STRONG_SELL"] if s in signals], [4.0, 3.0, 4.5, 4.5], align="lrrr")
    h2("11.6 Backtest")
    p(f"""A backtest converts forecasts into money. The rule is simple and fixed in advance: {BT['rule']}. Because a 21-day rebalancing
      schedule can start on any of 21 different days, and the result should not depend on that accident, the backtest is repeated for
      all {BH['runs']} start days and the median and range are reported ([[tab:backtest]]). [[fig:backtest]] plots the first of those
      runs.""")
    table("backtest", f"Backtest over {BH['runs']} start days, after 20 bps costs (top-5 strategy against all stocks equally weighted)",
          ["Measure", "Hold-out (2025 to 2026)", "Cross-validation years (2019 to 2024)"], [
              ["Strategy annual return, median", spct(BH["strategy_annual_return_median"]), spct(BC["strategy_annual_return_median"])],
              ["Benchmark annual return, median", spct(BH["benchmark_annual_return_median"]), spct(BC["benchmark_annual_return_median"])],
              ["Annual excess return, median", f"**{spct(BH['annual_excess_median'])}**", spct(BC["annual_excess_median"])],
              ["Annual excess return, range", f"{spct(BH['annual_excess_min'])} to {spct(BH['annual_excess_max'])}", f"{spct(BC['annual_excess_min'])} to {spct(BC['annual_excess_max'])}"],
              ["Start days on which the strategy won", f"{pct(BH['share_of_runs_beating_benchmark'], 0)}", f"{pct(BC['share_of_runs_beating_benchmark'], 0)}"],
              ["Maximum drawdown, median: strategy / benchmark", f"{pct(BH['strategy_max_drawdown_median'])} / {pct(BH['benchmark_max_drawdown_median'])}",
               f"{pct(BC['strategy_max_drawdown_median'])} / {pct(BC['benchmark_max_drawdown_median'])}"],
              ["First run: t-statistic of the per-period excess", f"{fh['excess_t_stat']:+.2f} ({fh['periods']} periods)", f"{fc['excess_t_stat']:+.2f} ({fc['periods']} periods)"],
              ["First run: probability the true excess is positive", pct(fh["prob_true_excess_positive"], 0), pct(fc["prob_true_excess_positive"], 0)],
          ], [6.6, 4.5, 4.9], align="lrr")
    figure("backtest", f"""Growth of one rupee in the first of the {BH['runs']} backtest runs. Left: the cross-validation years, using
      out-of-fold forecasts. Right: the hold-out period. Each panel has its own scale.""")
    p(f"""**The top-5 strategy did not beat the benchmark on the hold-out.** Its median annual return was
      {spct(BH['strategy_annual_return_median'])} against {spct(BH['benchmark_annual_return_median'])} for holding all
      {D['tickers']} stocks equally, a median shortfall of {pct(abs(BH['annual_excess_median']))} a year; it won from only
      {pct(BH['share_of_runs_beating_benchmark'], 0)} of start days, and its worst peak-to-trough fall was about three times the
      benchmark's ({pct(BH['strategy_max_drawdown_median'])} against {pct(BH['benchmark_max_drawdown_median'])}), which is what
      concentrating in five names does. In the cross-validation years the strategy was ahead by a median
      {pct(BC['annual_excess_median'])} a year, but the t-statistic of {fc['excess_t_stat']:.2f} is far below the conventional
      threshold of 2, and the probability that the true excess is positive (the probabilistic Sharpe ratio of Bailey and López de
      Prado, 2012) is {pct(fc['prob_true_excess_positive'], 0)}: not evidence of skill. The backtest also ignores taxes, market
      impact and any slippage beyond the flat 20 bps, all of which would make the strategy's result worse.""")
    h2("11.7 What the product can and cannot claim")
    callout("An honest reading of the results", f"""**Can claim:** (1) calibrated risk ranges: {pct(LG['coverage'])} hold-out coverage
      against an 80% target, with a band {pct(band_cut)} narrower than a no-feature baseline; (2) a properly validated pipeline in
      which tuning demonstrably repaired an over-fitted model ({pct(SEL['pinball_improvement_vs_notebook'])} lower loss);
      (3) consistent, explainable, profile-aware screening that saves time. **Cannot claim:** (1) that the model predicts returns
      better than historical percentiles (a tie in cross-validation, {pct(hold_gain)} better on the hold-out); (2) reliable stock-picking
      skill (rank IC unstable across periods); (3) that following its top picks beats holding the universe (it did not on the
      hold-out).""")
    p("""Why ship a model that only ties a naive baseline? Because the product is not sold on the forecast's edge. The naive baseline
      cannot tell a user that a real-estate stock has a wider range than a bank, cannot adapt its band when market volatility rises,
      and cannot explain itself. The quantile model does those things at no cost in accuracy. And a result of "no detectable edge",
      obtained with a leak-free method, is itself useful to an adviser: it is a documented reason not to promise clients
      market-beating picks.""")

    # ------------------------------------------------------------------ 12 SHAP
    h1("12. Explainability")
    h2("12.1 What SHAP values are")
    p("""A forecast that cannot be explained cannot be defended to a client. **SHAP** (SHapley Additive exPlanations; Lundberg and Lee,
      2017) explains one prediction at a time by sharing it out among the features. The idea comes from Shapley values in co-operative
      game theory: treat the features as players in a team, the prediction as the team's payout, and give each player the average of
      what it adds across all the orders in which the team could be assembled. The result has an exact accounting property: the
      model's average prediction (the *base value*) plus the SHAP values of all features equals the prediction for that row. A SHAP
      value of +0.004 for a feature means that this feature's value, for this stock today, moved the forecast 0.4 percentage points
      above where it would otherwise be. For tree models the `TreeExplainer` algorithm (Lundberg et al., 2020) computes these values
      exactly and quickly, which is one reason a tree model was chosen. LIME, the alternative named in the brief, fits a local
      approximation and gives slightly different answers on each run; SHAP's exactness and additivity made it the better fit.""")
    h2("12.2 Global importance")
    p(f"""Averaging the absolute SHAP values over a random sample of 5,000 hold-out rows gives each feature's overall influence on the
      median (P50) model ([[fig:shap]]). The largest driver is `{shap[0]['feature']}`, the market-wide volatility, with a mean absolute
      contribution of {shap[0]['mean_abs_shap'] * 100:.2f} percentage points; it is followed by `{shap[1]['feature']}`,
      `{shap[2]['feature']}`, `{shap[3]['feature']}` and `{shap[4]['feature']}`. The remaining five features together contribute
      {sum(s['mean_abs_shap'] for s in shap[5:]) * 100:.3f} percentage points: in practice the model uses five inputs.""")
    figure("shap", """Global SHAP importance of the median model on hold-out rows. Colour shows the direction of the relationship (the
      sign of the rank correlation between the feature's value and its SHAP value).""", 13.0)
    p(f"""The directions make economic sense and also explain the results of Section 11. Higher market volatility raises the forecast
      (direction {shap[0]['direction']:+.2f}), consistent with markets recovering after turbulent spells. Higher six-month momentum,
      a higher fractionally differenced price and a higher three-month return all *lower* the forecast (directions
      {shap[1]['direction']:+.2f}, {shap[2]['direction']:+.2f} and {shap[4]['direction']:+.2f}): the model has learned mean reversion,
      expecting stocks that have run up to give some back. The important point is the first one. The strongest feature is the same
      for every stock on a given day, so it moves all forecasts up or down together. It helps to set the *level and width* of the
      forecast band, which is where the model is good, but it cannot help to *rank* one stock above another, which is where the model
      is weak.""")
    h2("12.3 How one stock's explanation reads in the application")
    p("""On the Stock research tab the application runs `TreeExplainer` on the median model for the selected stock's latest row and lists
      up to eight features, largest first, each with its impact and whether it pushed the forecast up or down. An adviser reads it as a
      sentence. If, for example, the top rows were `mkt_vol_21` positive, `mom_126_21` negative and `fracdiff_close` negative, the
      explanation to the client would be: "the model's central estimate is a little above average because markets have been volatile,
      which has historically been followed by recovery, but it is held back because this share has already risen strongly over six
      months." (That example is illustrative wording, not a quoted output.) The same tab shows the technical indicators and the
      fundamental snapshot beside the SHAP list, so the adviser can check the model's reasons against the chart.""")

    # ------------------------------------------------------- 13 error analysis
    h1("13. Error analysis")
    p("""An average hides where a model fails. The hold-out rows were therefore split three ways: by market volatility (thirds of
      `mkt_vol_21` within the hold-out, labelled calm, normal and turbulent), by what the market then did (whether the average stock
      rose or fell over the following 21 days), and by sector. The second split uses the outcome, so it is a diagnosis and not
      something a user could act on in advance.""")
    table("regime", "Hold-out forecast quality by volatility regime and market direction", ["Group", "Rows", "MAE", "Coverage", "Band width", "Direction right", "Rank IC"],
          [[g["group"].capitalize(), f"{g['rows']:,}", f4(g["mae"]), pct(g["coverage"]), f"{g['band_width'] * 100:.1f} pp", pct(g["directional_accuracy"]), sic(g["ic_mean"])]
           for g in ERR["volatility_regime"] + ERR["market_direction"]], [3.2, 1.9, 1.9, 2.1, 2.4, 2.6, 1.9], align="lrrrrrr")
    figure("regime", """Hold-out error, coverage and rank IC by volatility regime (left three bars) and by the market's subsequent
      direction (right two bars). Each panel has its own scale.""")
    h2("13.1 By volatility regime")
    p(f"""Error rises with turbulence, as it should: the mean absolute error is {f4(regime['calm']['mae'])} in calm markets and
      {f4(regime['turbulent']['mae'])} in turbulent ones ([[tab:regime]], [[fig:regime]]). The model responds correctly by widening
      its band from {regime['calm']['band_width'] * 100:.1f} to {regime['turbulent']['band_width'] * 100:.1f} percentage points, and
      coverage stays within about five points of target in every regime ({pct(regime['calm']['coverage'])},
      {pct(regime['normal']['coverage'])}, {pct(regime['turbulent']['coverage'])}). This is the behaviour a risk tool needs. Ranking is
      a different matter: rank IC is {sic(regime['turbulent']['ic_mean'])} in turbulent markets but *negative*
      ({sic(regime['calm']['ic_mean'])}) in calm ones, where it was positive on only {pct(regime['calm']['ic_positive_share'], 0)} of
      days. The mean-reversion pattern the model learned pays when prices are moving violently and works against it when quiet trends
      persist.""")
    h2("13.2 By market direction")
    p(f"""When the market went on to fall, the median forecast had the direction right only
      {pct(direction['market fell']['directional_accuracy'])} of the time and coverage dropped to
      {pct(direction['market fell']['coverage'])}; when it rose, the figures were {pct(direction['market rose']['directional_accuracy'])}
      and {pct(direction['market rose']['coverage'])}. The model does not forecast market-wide falls. Its median is usually slightly
      positive, reflecting the long-run upward drift in its training data, so it is "right" in rising months and "wrong" in falling
      ones. A user must not read a positive median as protection against a market decline; the P10 figure is the number that speaks to
      that risk, and even it is breached more often than intended when the whole market falls.""")
    h2("13.3 By sector")
    table("sector", "Hold-out forecast quality by sector, ordered by coverage", ["Sector", "MAE", "Coverage", "Band width", "Direction right", "Rank IC"],
          [[g["group"], f4(g["mae"]), pct(g["coverage"]), f"{g['band_width'] * 100:.1f} pp", pct(g["directional_accuracy"]), sic(g["ic_mean"])] for g in sectors],
          [5.6, 1.8, 2.0, 2.2, 2.6, 1.8], align="lrrrrr")
    figure("sector", "Hold-out coverage and error by sector. Sectors are ordered by coverage, lowest at the bottom.")
    p(f"""Coverage is clearly below target in three sectors: {sectors[0]['group']} ({pct(sectors[0]['coverage'])}), {sectors[1]['group']}
      ({pct(sectors[1]['coverage'])}) and {sectors[2]['group']} ({pct(sectors[2]['coverage'])}); these are also the sectors with the
      largest errors ([[tab:sector]], [[fig:sector]]). It is above target in defensive sectors such as {sectors[-1]['group']}
      ({pct(sectors[-1]['coverage'])}) and {sectors[-2]['group']} ({pct(sectors[-2]['coverage'])}). The reason is structural: the model
      has no sector input and its strongest feature is market-wide, so it gives broadly similar band widths to all stocks (between
      {min(g['band_width'] for g in sectors) * 100:.1f} and {max(g['band_width'] for g in sectors) * 100:.1f} percentage points) although
      their true ranges differ by more. Bands are a little too narrow for volatile sectors and a little too wide for stable ones. The
      sector rank IC values should be read with caution, since each is a correlation across only five stocks a day. Adding sector, or
      calibrating bands sector by sector, is the first modelling improvement we would make.""")

    # ------------------------------------------------------ 14 decision engine
    h1("14. The decision engine")
    p("""The forecast is one input to the recommendation, not the recommendation itself. Given how modest the forecast's edge is, this
      separation is a design decision and not a detail. The decision engine (`src/portfolio_advisor/core.py`) is deterministic,
      rule-based code: the same inputs always give the same output, and every rule can be read.""")
    h2("14.1 Four pillars")
    p("""For every stock on the chosen date the engine computes four scores, each a z-score across the universe (how many standard
      deviations the stock is from the day's average), clipped to ±3.""")
    bullets([
        """**Technical.** The weighted average of eleven indicator z-scores. RSI, MACD, 5-day and 21-day return, fractionally
        differenced price, on-balance-volume change and stochastic %K count in favour; Williams %R, 21-day volatility, ATR and
        Bollinger width count against. The 21-day return, MACD and differenced price carry a weight of 1.25, the rest 1. (Because
        Williams %R is stochastic %K shifted by 100, those two terms offset each other in the current code, so the score effectively
        rests on the other nine.)""",
        """**Fundamental.** The weighted average of seven ratio z-scores from today's snapshot. Low price-to-earnings, low price-to-book
        and low debt-to-equity score well, as do high return on equity, net margin, dividend yield and revenue growth. Return on
        equity, net margin and revenue growth carry a weight of 1.15, the rest 1. Valuation ratios are log-transformed first so that one extreme value does not dominate.""",
        """**Model.** The z-score of the stock's median 21-day forecast among that day's forecasts.""",
        """**Portfolio fit (risk).** −0.55 × the z-score of trailing annual volatility − 0.45 × the z-score of trailing maximum
        drawdown, so that calmer stocks score higher.""",
    ])
    h2("14.2 Combining the pillars by risk profile")
    p("""The overall score is a weighted sum of the four pillars, and the weights are what make the product profile-aware
      ([[tab:presets]]). A Conservative profile leans on fundamentals and risk and gives the model forecast the least say (15%); an
      Aggressive profile leans on technical momentum and the model. The weighted sum is then mapped to a 0 to 100 scale by
      50 + 50 × tanh(x ÷ 1.5), which keeps an average stock at 50 and compresses extremes. The same table shows the limits each profile
      applies before a stock may enter the portfolio.""")
    sys.path.insert(0, str(ROOT / "src"))
    from portfolio_advisor.core import PROFILE_PRESETS  # noqa: PLC0415

    presets = list(PROFILE_PRESETS.values())
    table("presets", "Risk-profile presets from `core.py`: pillar weights and eligibility limits", ["Setting"] + [x.name for x in presets], [
        ["Weight: technical"] + [pct(x.technical_weight, 0) for x in presets],
        ["Weight: fundamental"] + [pct(x.fundamental_weight, 0) for x in presets],
        ["Weight: model forecast"] + [pct(x.model_weight, 0) for x in presets],
        ["Weight: portfolio fit (risk)"] + [pct(x.risk_weight, 0) for x in presets],
        ["Minimum confidence"] + [f"{x.min_confidence:.2f}" for x in presets],
        ["Maximum annual volatility"] + [pct(x.max_annual_volatility, 0) for x in presets],
        ["Maximum trailing drawdown"] + [pct(x.max_drawdown, 0) for x in presets],
        ["Maximum weight of one position"] + [pct(x.max_position_weight, 0) for x in presets],
        ["Minimum fundamental score (0 to 100)"] + [f"{x.min_fundamental_score:.0f}" for x in presets],
    ], [6.4, 3.2, 3.2, 3.2], align="lrrr")
    h2("14.3 Recommendation states")
    p("""Two further quantities enter the rating. *Confidence* is one minus the percentile rank of the stock's P10 to P90 band width
      among that day's bands: the narrowest band has confidence near 1. *Evidence confidence* blends it with agreement between pillars:
      0.60 × confidence + 0.25 × agreement (1 if the technical and fundamental scores have the same sign, otherwise 0.5) + 0.15 if
      fundamental data exist for the stock. A stock is *model-positive* if its model signal is a buy or its median forecast is not
      negative. The four states follow in order ([[tab:states]]).""")
    table("states", "Recommendation states", ["Rating", "Rule"], [
        ["STRONG BUY", "Model-positive, overall score at least 70, and evidence confidence at least the profile's minimum confidence"],
        ["BUY", "Model-positive, overall score at least 50, and evidence confidence at least 85% of the profile's minimum"],
        ["PASS", "Not model-positive, or overall score below 42"],
        ["WATCH", "Everything else: evidence is mixed and needs confirmation"],
    ], [3.0, 13.0])
    p("""Only STRONG BUY and BUY stocks can enter the portfolio, and only if they also pass every profile limit: confidence, minimum
      expected return, volatility, drawdown, fundamental score, and any sector or ticker exclusion. Eligible stocks are ranked by
      overall score and the top *N* are taken, where *N* is the user's maximum number of holdings. If nothing qualifies, the engine
      shows the best WATCH names, clearly marked as a watchlist, instead of forcing a portfolio. Each row also carries plain-language
      risk flags (volatility above 60%, drawdown above 40%, price-to-earnings above 40, debt-to-equity above 150, or confidence below
      0.35) and a one-sentence reason.""")
    h2("14.4 Position sizing: Hierarchical Risk Parity, step by step")
    p("""Having chosen the stocks, the engine must decide how much of each to hold. Equal weights ignore risk. Classical mean-variance
      optimisation needs expected returns, which Section 11 shows we cannot estimate reliably, and it requires inverting the
      covariance matrix, which makes its weights unstable. **Hierarchical Risk Parity** (HRP; López de Prado, 2016) uses only risks
      and correlations and needs no matrix inversion. The engine applies it to the last 252 daily returns of the selected stocks.""")
    numbered([
        """**Measure similarity.** Compute the correlation between every pair of stocks and convert it to a distance,
        √(0.5 × (1 − correlation)), so that stocks moving together are close and stocks moving independently are far apart.""",
        """**Cluster.** Build a tree by repeatedly joining the two closest stocks or groups (single-linkage hierarchical clustering).
        Stocks from the same sector typically join early.""",
        """**Reorder.** List the stocks in the order of the tree's leaves, so that similar stocks sit next to each other
        (quasi-diagonalisation).""",
        """**Split the capital top-down.** Cut the ordered list in half. Estimate each half's variance using inverse-variance weights
        inside it, and give each half a share of capital inversely proportional to its variance: the riskier half gets less. Repeat
        inside each half until every stock stands alone (recursive bisection).""",
    ])
    p("""The effect is that two highly correlated bank stocks share one allocation between them instead of each taking a full slice, and
      a volatile stock receives less capital than a calm one. If fewer than 30 days of common history exist, or the clustering fails,
      the engine falls back to equal weights instead of guessing.""")
    h2("14.5 Tilt, position cap, whole shares and cash")
    p("""Four practical steps follow. (1) **Tilt:** each HRP weight is multiplied by the stock's overall score ÷ 100 (limited to between
      0.25 and 1.25) and the weights are rescaled to sum to one, so better-supported names get somewhat more. (2) **Cap:** no position
      may exceed the profile's maximum weight (20%, 30% or 40%); any excess is redistributed among the uncapped positions, repeatedly,
      until all comply. If the cap multiplied by the number of positions is below 100%, the remainder stays in cash. (3) **Whole
      shares:** the number of shares is the rupee allocation divided by the latest price, rounded *down*, because Indian cash equities
      trade in whole shares and the user must never be asked to spend more than the stated capital. (4) **Cash remainder:** whatever
      rounding and the cap leave over is reported as cash, and the summary shows the weights actually achieved. A unit test checks that
      the amount invested never exceeds the capital. The Portfolio lab tab also shows an illustrative schedule that splits each order
      into ten decreasing slices; it is a display only and no order is sent anywhere.""")

    # -------------------------------------------------------- 15 engineering
    h1("15. Engineering and reproducibility")
    h2("15.1 Reproducibility in three commands")
    code(["pip install -r requirements.txt", "python -m src.model          # select, tune, cross-validate, train, test, export",
          "streamlit run app.py         # open the product at http://localhost:8501"])
    p(f"""The second command retrains everything from the committed data snapshot in about a minute on a laptop and rewrites
      `reports/metrics.json`, `reports/cv_results.csv` and the model bundle. All randomness (the feature-selection sample, the random
      search, row and feature sampling inside LightGBM, the SHAP sample) uses the fixed seed {CFG['SEED']}. A fourth command,
      `python scripts/build_report.py`, regenerates this report and every figure in it from the two files in `reports/`, so no number
      in the text is typed by hand. `python -m src.data --refresh` downloads fresh prices when an update is wanted.""")
    h2("15.2 Automated tests")
    p("""The repository has 22 automated tests, run with `pytest` and on every push by a GitHub Actions workflow that also
      byte-compiles the sources ([[tab:tests]]). The six pipeline tests are the ones that matter most for a grader: they turn the
      leakage controls of Section 9 from statements into checks that fail the build if someone breaks them.""")
    table("tests", "The 22 automated tests", ["File", "Tests", "What is checked"], [
        ["`tests/test_pipeline.py`", "6", "Walk-forward folds never train on the future and keep the purge gap; the hold-out is separated from development by the gap; no fundamental feature is a model input; fractional differencing with d = 1 equals the ordinary difference; pinball-loss values and the P10 ≤ P50 ≤ P90 ordering after crossing; the saved report is internally consistent"],
        ["`tests/test_core.py`", "5", "All four pillar scores are present; ticker and sector exclusions are obeyed; invested amount never exceeds capital; the order schedule sums to the target whole shares; validation and data-quality reports are available; outputs convert cleanly to JSON"],
        ["`tests/test_chat.py`", "5", "The chat context covers the whole universe; the grounding prompt contains the data table; success, retry-on-transient-failure and not-configured paths"],
        ["`tests/test_live.py`", "3", "Intraday return uses the previous close; the default provider is explicitly labelled not real-time; an Upstox feed message decodes correctly without a network"],
        ["`tests/test_api.py`", "2", "The recommendations endpoint returns the profile analysis; the health endpoint reports a degraded state when the model bundle is missing"],
        ["`tests/test_app.py`", "1", "The Streamlit application renders every tab against the real model bundle without an exception"],
    ], [3.8, 1.2, 11.0], align="lrl", size=8)
    h2("15.3 Deployment")
    p(f"""The product is packaged as a Docker image built from `python:3.12-slim`. The image installs the pinned requirements, copies
      the source, the application and the `reports/` folder, exposes port 8501 and declares a health check against Streamlit's
      `/_stcore/health` endpoint. The model bundle is mounted read-only at `/app/data`, so the public container can read forecasts but
      cannot alter them. A `docker-compose.yml` file adds the optional FastAPI service as a second container. The live instance at
      {APP_URL} runs on Coolify, an open-source self-hosted deployment platform, on a virtual private server; a separate landing page
      at {LANDING_URL} introduces the product. Secrets (broker tokens, the chat API key) are supplied as environment variables in the
      platform and are excluded from the repository by `.gitignore`.""")
    h2("15.4 Vibe-coding workflow")
    p("""The course asks teams to build with AI coding agents and to document how. The full record is `docs/VIBE_CODING_LOG.md`; this is
      a summary. The platform was Claude Code, Anthropic's terminal coding agent, with one orchestrating session and a background
      sub-agent, supported by Playwright for browser screenshots and by `uv`, `pytest` and the GitHub command line.""")
    p("""The work followed four steps. **Audit before building:** the agent was given the assignment and asked which existing project met
      it; it scored the project against the rubric row by row and found that the application and SHAP existed but tuning,
      cross-validation, the required file layout and the log did not. **Gap-driven build:** each missing rubric row became a concrete
      change, namely the random search, purged walk-forward validation, a scaled Ridge baseline, removal of the leaking fundamentals,
      and the move from a notebook to `src/data.py` and `src/model.py`. **Parallel delegation:** the landing page went to a sub-agent
      with a written brief stating its folder, the facts it could use, and a rule that no performance number could appear until real
      metrics existed. **Verify, then report:** every change was run end to end, including the training script, the tests and a render
      of the application.""")
    p("""Four prompting practices worked: giving the agent the grading rubric instead of a task list, so it could rank work by marks;
      asking for evidence (a met-or-missing table with file references) instead of a yes or no; briefing sub-agents like a colleague,
      with scope, facts and forbidden claims; and demanding honest numbers, which is how the tie with the naive baseline and the losing
      backtest came to be reported unchanged. Three things went wrong and are recorded: an ambiguous "this" sent the agent to the wrong
      repository for about an hour; joining forecasts onto a date index that repeats once per stock silently multiplied rows until a
      length check exposed it (forecasts are now attached by position, with a comment explaining why); and the notebook's settings
      turned out to be over-complex, which only cross-validation revealed.""")
    p("""The course's golden rule is that AI may write the scaffolding but the team must be able to explain every algorithm and
      parameter. Appendix A is the team's preparation for that test.""")

    # ------------------------------------------------------------ 16 limits
    h1("16. Limitations, risks, ethics and next steps")
    h2("16.1 Limitations")
    bullets([
        """**The forecast has little or no edge.** Returns over 21 days are mostly noise. The model matches, and does not beat,
        historical percentiles in cross-validation; its ranking skill is unproven; its top picks lost to the benchmark on the hold-out.""",
        f"""**Survivorship bias and a narrow universe.** {D['tickers']} of today's large caps in one country. Results may not carry
        over to mid caps, small caps or other markets.""",
        """**One hold-out period.** About 20 months of a single market regime. A different period could give different numbers in
        either direction.""",
        """**Fundamentals are a snapshot.** The fundamental pillar uses today's ratios. It is valid for today's decision but has never
        been validated historically, because point-in-time fundamentals were not available.""",
        """**The decision engine's weights are judgement.** The pillar weights, thresholds (70, 50, 42) and HRP tilt were set by
        design reasoning, not fitted or backtested. Only the forecast model has been validated.""",
        """**Calibration varies by sector and period.** Bands are too narrow for IT, real estate and consumer discretionary stocks and
        when the whole market falls.""",
        """**Simplified costs.** The backtest charges a flat 20 bps and ignores taxes, market impact and liquidity.""",
        """**Cross-validation scores are slightly optimistic** because feature selection saw the whole development set (Section 9.5).""",
    ])
    h2("16.2 Risks in use")
    p("""The main risk is **misplaced trust**: a user reading "STRONG BUY" as a promise. The mitigations are in the product: ranges
      instead of points, a Validation tab and Model report tab that show the weak results openly, a warning under the backtest, and
      research-only wording on every screen. A second risk is **staleness**: the model is retrained on demand, not continuously, so an
      old bundle could be served; the application prints the artefact date on every page. A third is **data quality**: Yahoo Finance is
      a free source without a service guarantee, and an erroneous price would flow into the features. A fourth is the **chat
      assistant**, which is a language model and can misstate figures even when given the right table; it is labelled research-only and
      is switched off unless configured.""")
    h2("16.3 Ethics and compliance")
    p("""AI Portfolio Advisor is a **research and education tool**. It is not investment advice, it does not know a user's full
      financial situation, and its authors are not registered investment advisers or research analysts. In India, giving investment
      advice for a fee and publishing research recommendations are regulated by the Securities and Exchange Board of India (SEBI)
      under its Investment Advisers and Research Analysts regulations. A commercial launch would therefore need legal review, and the
      natural customer is an adviser who is already registered and who remains responsible for the advice given. The product supports
      that responsibility by making every recommendation traceable to its evidence. It stores no personal data: the profile is held only
      for the duration of the session. And it places no trades.""")
    p(f"""We also regard honest reporting as an ethical matter. A financial product that advertised the cross-validation backtest
      ({spct(BC['annual_excess_median'])} a year) and omitted the hold-out result ({spct(BH['annual_excess_median'])}) would mislead. Both are shown here and in the application.""")
    h2("16.4 Next steps")
    bullets([
        """**Point-in-time fundamentals,** so that the fundamental pillar can be validated and, if it earns its place, added to the model without leakage.""",
        """**A market-relative label** (stock return minus the market's return), which targets ranking directly and removes the market-wide component the model cannot predict.""",
        """**Sector-aware calibration** of the bands, for example conformal adjustment of P10 and P90 by sector.""",
        """**A wider, survivorship-free universe** that includes companies later delisted or demoted.""",
        """**Nested cross-validation,** with feature selection repeated inside each fold.""",
        """**A paper-trading period** of at least six months, and a time-and-motion pilot with advisers to replace the ROI assumptions with measurements.""",
    ])

    # ------------------------------------------------------------ references
    h1("References")
    for ref in [
        "Bailey, D. H. and López de Prado, M. (2012). The Sharpe ratio efficient frontier. *Journal of Risk*, 15(2).",
        "Bergstra, J. and Bengio, Y. (2012). Random search for hyper-parameter optimization. *Journal of Machine Learning Research*, 13, 281–305.",
        "Dickey, D. A. and Fuller, W. A. (1979). Distribution of the estimators for autoregressive time series with a unit root. *Journal of the American Statistical Association*, 74(366), 427–431.",
        "Friedman, J. H. (2001). Greedy function approximation: a gradient boosting machine. *Annals of Statistics*, 29(5), 1189–1232.",
        "Hoerl, A. E. and Kennard, R. W. (1970). Ridge regression: biased estimation for nonorthogonal problems. *Technometrics*, 12(1), 55–67.",
        "Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q. and Liu, T.-Y. (2017). LightGBM: a highly efficient gradient boosting decision tree. *Advances in Neural Information Processing Systems*, 30.",
        "Koenker, R. and Bassett, G. (1978). Regression quantiles. *Econometrica*, 46(1), 33–50.",
        "López de Prado, M. (2016). Building diversified portfolios that outperform out of sample. *Journal of Portfolio Management*, 42(4), 59–69.",
        "López de Prado, M. (2018). *Advances in Financial Machine Learning*. Wiley.",
        "Lundberg, S. M. and Lee, S.-I. (2017). A unified approach to interpreting model predictions. *Advances in Neural Information Processing Systems*, 30.",
        "Lundberg, S. M., Erion, G., Chen, H., et al. (2020). From local explanations to global understanding with explainable AI for trees. *Nature Machine Intelligence*, 2, 56–67.",
        "Peng, H., Long, F. and Ding, C. (2005). Feature selection based on mutual information: criteria of max-dependency, max-relevance, and min-redundancy. *IEEE Transactions on Pattern Analysis and Machine Intelligence*, 27(8), 1226–1238.",
    ]:
        _add("ref", text=ref)

    # ------------------------------------------------------------ appendix A
    h1("Appendix A. Technical-defence questions and answers", page_break=True)
    p("""These are questions an examiner could put to any member of the team. Each answer is grounded in this project's code and
      results, and every team member should be able to give it without notes.""")
    h2("A.1 Problem and data")
    qa("What exactly is the model predicting?", f"""The percentage change in a stock's adjusted closing price over the next
      {CFG['HORIZON']} trading days. It predicts three percentiles of that return, the 10th, 50th and 90th, and not a single number,
      so the output is a range. It is a regression problem, not a classification problem.""")
    qa("Why 21 trading days?", """It is about one calendar month, which matches how often our target user reviews a portfolio. A shorter
      horizon would be noisier and would imply more trading than an adviser's client does; a longer one would leave too few
      non-overlapping periods in eight years of data to validate anything.""")
    qa("What is survivorship bias and how does it affect you?", """We picked 60 companies that are large today, so companies that failed
      or shrank since 2018 are missing. That inflates the historical return of the universe. We limit the damage by comparing the
      strategy with the same 60 stocks held equally, so both sides carry the same bias, and we state it as a limitation.""")
    qa("Why did you remove fundamentals from the model when they obviously matter?", """Because we only have today's values. Copying a
      2026 price-to-earnings ratio onto a 2019 row tells the model something about the future, which is look-ahead leakage and inflates
      test scores. Fundamentals still feed the application's separate fundamental score for today's decision, where no look-ahead
      exists, and a unit test stops them returning to the model.""")
    h2("A.2 Features")
    qa("What is fractional differencing and why not just use returns?", """Returns are stationary but forget almost everything about
      where the price has been. Fractional differencing subtracts a slowly decaying weighted sum of past prices, using an order d
      between 0 and 1, which removes the trend while keeping long memory. We pick the smallest d that passes the ADF stationarity test,
      per stock, using development data only.""")
    qa("What does the ADF test tell you?", """Its null hypothesis is that the series has a unit root, meaning it wanders with no fixed
      mean. A p-value below 0.05 rejects that, so we treat the series as stationary. We use it as a stopping rule for choosing d, not as
      proof of anything about returns.""")
    qa("How does your feature selection work, and why not just use correlation?", f"""It is greedy: take the feature with the highest
      mutual information with the label, then repeatedly add the feature with the best relevance minus half its average mutual
      information with those already chosen, until there are {n_selected}. Mutual information also picks up non-linear relationships
      that correlation misses, and the redundancy term stops us keeping near-duplicates such as Stoch_K and Williams_R.""")
    qa("One selected feature has zero relevance. Is that a bug?", """No, it is a known weakness of the method that we report. OBV_delta
      was chosen ninth because, once the informative features are in, a feature that is merely different scores better than one that
      is informative but redundant. SHAP confirms it contributes almost nothing, so it does no harm, and a small tree model simply
      ignores it.""")
    h2("A.3 Model and loss")
    qa("Why pinball loss and not mean squared error?", """Minimising squared error gives the mean, a single number. We need percentiles
      to show a range, and the loss whose minimiser is the q-th percentile is the pinball loss, an absolute error with slopes q and
      1 − q on the two sides. It is also less sensitive to extreme returns than squared error.""")
    qa("Why three separate models instead of one?", """Each percentile has a different loss function, and a LightGBM model minimises one
      loss. So there is one model each for q = 0.1, 0.5 and 0.9, with identical features and settings. Because they are independent
      their outputs can occasionally cross, so we sort the three predictions for every row.""")
    qa("Explain gradient boosting in two sentences.", """Start with a constant prediction, then repeatedly fit a small decision tree to
      the errors of the current model and add a fraction of it. The fraction is the learning rate, and the final model is the sum of
      all the small trees.""")
    qa(f"What does num_leaves = {params['num_leaves']:g} do?", f"""It limits every tree to {params['num_leaves']:g} end-points, which
      means at most three splits and at most three features interacting in any one tree. The notebook used 31. The search preferred the
      smaller value because larger trees fitted noise in past returns and scored worse on later periods.""")
    qa(f"Why {params['n_estimators']:g} trees with a learning rate of {params['learning_rate']:g}?", """Together they set how far the
      model can move from its starting point, the historical percentile. One hundred trees each contributing 2% of their proposed
      correction is a small total adjustment. No setting with 400 or more trees ranked in the top twelve.""")
    qa("What do subsample and colsample_bytree do?", f"""`subsample` = {params['subsample']:g} gives each tree a random 80% of the rows
      and `colsample_bytree` = {params['colsample_bytree']:g} a random half of the features. Trees built on different samples make
      different mistakes, and averaging them reduces variance.""")
    qa("Why did you scale features for Ridge but not for LightGBM?", """Ridge penalises coefficient size, which depends on each
      feature's units, so features must be standardised for the penalty to be fair; the scaler is refitted on each training fold to
      avoid leakage. A tree only asks whether a value is above a threshold, so rescaling changes nothing.""")
    qa("Why LightGBM and not a neural network or an LSTM?", """We have about 91,000 development rows of tabular data with a very weak
      signal. Gradient-boosted trees are the standard strong choice for tabular data, train in seconds, support the quantile loss
      directly and have exact SHAP explanations. Our own results show that even a 31-leaf tree over-fits here, so a far more flexible
      model would be expected to do worse.""")
    h2("A.4 Validation and tuning")
    qa("Why not ordinary k-fold cross-validation?", """It shuffles rows, so it trains on the future to predict the past, and neighbouring
      rows share most of their 21-day label window, so a shuffled test set contains near-copies of training rows. Both make scores look
      better than live use. Walk-forward validation always trains on the past and tests on the next block.""")
    qa(f"Why purge {GAP} days?", f"""A training label at day t uses the price at t + {CFG['HORIZON']}. Without a gap, the last
      {CFG['HORIZON']} training labels would be computed from prices inside the validation period. We drop those {CFG['HORIZON']} days
      (the purge) and a further {CFG['EMBARGO']} (the embargo) because market conditions persist for days.""")
    qa("What is the difference between the validation folds and the hold-out?", """The folds were used to choose among 25 settings, so
      their best score is optimistic. The hold-out, the latest 20% of dates, was used for nothing until the end and was scored once. It
      is the estimate of live performance.""")
    qa("Why random search and not grid search or Bayesian optimisation?", f"""The grid has
      {int(np.prod([len(v) for v in M['validation']['search_space'].values()])):,} combinations and each needs 15 model fits. Random
      search samples it efficiently when only a few parameters matter, which is the case here. With a signal this weak, the differences
      between good settings are within noise, so a more elaborate optimiser would mostly be fitting the validation folds.""")
    qa("How do you know you did not over-fit the hyperparameters?", f"""Three ways. The chosen setting is one of the simplest in the
      space, not an exotic one. The top five settings are within {TRIALS.iloc[4]['mean'] - TRIALS.iloc[0]['mean']:.5f} of each other, so
      nothing was gained by fine selection. And the hold-out, which the search never saw, gave a result consistent with
      cross-validation.""")
    h2("A.5 Results and product")
    qa("Why is your model no better than the naive baseline, and why ship it?", f"""Because 21-day stock returns are close to
      unpredictable from past prices, and a leak-free validation shows that honestly: {f5(TUNED['pinball_mean'])} against
      {f5(NAIVE['pinball_mean'])} in cross-validation. We ship it because the product's value is calibrated ranges that adapt to
      market volatility, explanations and consistent screening, none of which the naive baseline provides, and we make no claim of
      superior returns.""")
    qa(f"What does {pct(LG['coverage'])} coverage mean, and is it good?", f"""On the hold-out, {pct(LG['coverage'])} of actual 21-day returns fell
      between the model's P10 and P90 forecasts. The band is designed to hold 80%, so it is slightly wide, which is the safe side. It
      varies by period and sector, so it is an average property and not a guarantee.""")
    qa("What is rank IC and what is yours?", f"""It is the daily Spearman rank correlation between forecasts and subsequent returns
      across the 60 stocks. Ours is {sic(LG['ic_mean'])} on the hold-out but {sic(TUNED['ic_mean'])} ± {TUNED['ic_std']:.3f} across
      cross-validation folds, so we do not claim stable ranking skill.""")
    qa("Your backtest lost to the benchmark. What do you conclude?", f"""That buying the five highest forecasts is not a strategy we can
      recommend: it returned {spct(BH['strategy_annual_return_median'])} a year against {spct(BH['benchmark_annual_return_median'])}
      for holding everything, with about three times the drawdown. It confirms that the forecast should be one input among four in the
      decision engine and should carry modest weight.""")
    qa("How would you explain a SHAP value to a client?", """Every forecast starts at the model's average. Each feature then pushes it up
      or down by an amount, and those amounts add up exactly to the forecast. A SHAP value of +0.4 points for market volatility means
      today's volatility raised this stock's median forecast by 0.4 percentage points.""")
    qa("What is HRP and why not mean-variance optimisation?", """Hierarchical Risk Parity clusters stocks by correlation and then splits
      capital down the cluster tree, giving less to riskier branches. Mean-variance optimisation needs expected returns, which we have
      shown we cannot forecast reliably, and inverts the covariance matrix, which makes weights unstable. HRP needs neither.""")
    qa("If the forecast is weak, what are the ratings based on?", """On four pillars combined with stated weights: technical indicators,
      today's fundamentals, the model forecast and trailing risk. For a Conservative profile the model forecast carries only 15% of the
      score. The ratings are a transparent screening rule, and we say plainly that the rule's weights were set by judgement and have
      not been backtested.""")
    qa("Is the ROI real?", f"""The arithmetic is real; the inputs are assumptions. {ROI['clients']} clients, {ROI['reviews_per_year']}
      reviews, {saved} minutes saved and Rs {ROI['adviser_rate_inr']:,} an hour give {pct(BASE['roi'], 0)}. We show a sensitivity
      table, identify a break-even of about {break_even:.0f} clients, and say that a time-and-motion pilot is needed to replace the
      assumptions with measurements.""")
    qa("Which parts did the AI agent write, and which do you own?", """The agent wrote most of the code under our direction and
      verified it by running it. We own the decisions: the problem, the choice of quantile forecasts, the insistence on purged
      validation and the removal of leaking features, the decision to report unfavourable results, and the ability to explain every
      parameter in this appendix.""")

    # ------------------------------------------------------------ appendix B
    h1("Appendix B. All hyperparameter trials")
    p(f"""All {len(TRIALS)} settings, ranked by mean pinball loss over the five folds. Trial 0 is the original notebook setting; trial
      {BEST_TRIAL} was chosen. The per-fold rows are in `reports/cv_results.csv`.""")
    table("trials_full", "Full hyperparameter search results", trial_header + ["Rank IC"],
          [trial_row(i + 1, r) + [f"{r['ic_mean']:+.3f}"] for i, (_, r) in enumerate(TRIALS.iterrows())],
          [0.7, 2.3, 1.7, 1.5, 1.1, 1.3, 1.4, 1.5, 1.1, 1.3, 0.8, 1.3], align="llrrrrrrrrrr", size=7.5)

    # ------------------------------------------------------------ appendix C
    h1("Appendix C. Glossary", page_break=True)
    table("glossary", "Glossary of technical terms", ["Term", "Meaning"], [[f"**{term}**", meaning] for term, meaning in GLOSSARY], [3.4, 12.6], size=8)


def check_narrative() -> list[str]:
    """Warn when a retrain has changed a result that the prose states in words."""
    checks = [
        (TUNED["pinball_mean"] < NOTEBOOK["pinball_mean"], "tuned model no longer beats the notebook settings in CV"),
        (abs(SEL["pinball_improvement_vs_naive"]) < 0.02, "tuned model no longer 'ties' the naive baseline in CV (gap above 2%)"),
        (LG["pinball"] < NV["pinball"], "tuned model no longer beats naive on the hold-out"),
        (0.78 <= LG["coverage"] <= 0.86, "hold-out coverage is no longer close to the 80% target"),
        (BH["annual_excess_median"] < 0, "hold-out backtest now beats the benchmark; Section 11.6 says it does not"),
        (M["shap"][0]["feature"] == "mkt_vol_21", "top SHAP feature changed; Section 12.2 names mkt_vol_21"),
        (SEL["params"]["num_leaves"] == 4 and SEL["params"]["n_estimators"] == 100, "chosen hyperparameters changed; Sections 10.3 and A.3 describe 4 leaves, 100 trees"),
        (RG["mae"] < LG["mae"], "Ridge no longer has the lowest hold-out MAE; see Section 11.2"),
    ]
    return [message for ok, message in checks if not ok]


def resolve_references() -> None:
    numbers: dict[str, str] = {}
    n_fig = n_tab = 0
    for block in BLOCKS:
        if block["kind"] == "figure":
            n_fig += 1
            block["number"] = n_fig
            numbers[f"fig:{block['key']}"] = f"Figure {n_fig}"
        elif block["kind"] == "table":
            n_tab += 1
            block["number"] = n_tab
            numbers[f"tab:{block['key']}"] = f"Table {n_tab}"

    def fix(value):
        if isinstance(value, str):
            return re.sub(r"\[\[(\w+:\w+)\]\]", lambda m: numbers[m.group(1)], value)
        if isinstance(value, list):
            return [fix(v) for v in value]
        return value

    for block in BLOCKS:
        for key, value in list(block.items()):
            if key != "key":
                block[key] = fix(value)


# ------------------------------------------------------------------ Markdown
def _md_cell(text: str) -> str:
    return text.replace("|", "\\|")


def render_markdown(figs: dict[str, Path]) -> str:
    out: list[str] = []
    for b in BLOCKS:
        kind = b["kind"]
        if kind == "title":
            out += ["# AI Portfolio Advisor", "", f"*{SUBTITLE}*", "", f"**Capstone project report** · {COURSE}", "",
                    f"**Team ({GROUP}):** " + "; ".join(f"{name} ({roll})" for name, roll in TEAM), "",
                    f"**Live app:** {APP_URL} · **Landing page:** {LANDING_URL} · **Code:** {REPO_URL}", "",
                    f"*Generated by `scripts/build_report.py` from `reports/metrics.json` ({M['generated']}). Word version: [Capstone_Report.docx](Capstone_Report.docx).*", "",
                    f"> {DISCLAIMER}", ""]
        elif kind == "toc":
            out += ["## Contents", ""] + [f"- {x['text']}" for x in BLOCKS if x["kind"] == "h1"] + [""]
        elif kind == "h1":
            out += [f"## {b['text']}", ""]
        elif kind == "h2":
            out += [f"### {b['text']}", ""]
        elif kind in ("p", "ref"):
            out += [b["text"], ""]
        elif kind == "bullets":
            out += [f"- {i}" for i in b["items"]] + [""]
        elif kind == "numbered":
            out += [f"{n}. {i}" for n, i in enumerate(b["items"], 1)] + [""]
        elif kind == "table":
            out += [f"**Table {b['number']}. {b['caption']}**", "", "| " + " | ".join(b["header"]) + " |",
                    "|" + "|".join("---:" if a == "r" else "---" for a in b["align"]) + "|"]
            out += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in b["rows"]] + [""]
        elif kind == "figure":
            out += [f"![Figure {b['number']}](figures/{figs[b['key']].name})", "", f"*Figure {b['number']}. {b['caption']}*", ""]
        elif kind == "callout":
            out += [f"> **{b['title']}.** {b['text']}", ""]
        elif kind == "eq":
            out += ["```", b["text"], "```", ""]
        elif kind == "code":
            out += ["```bash", *b["lines"], "```", ""]
        elif kind == "qa":
            out += [f"**Q. {b['q']}**", "", b["a"], ""]
    return "\n".join(out).rstrip() + "\n"


# ------------------------------------------------------------------- Word
# Cambria, Arial and Courier New have metric-compatible substitutes in LibreOffice (Caladea, Liberation Sans,
# Liberation Mono), so the page numbers read back from the PDF also hold when the file is opened in Word.
BODY_FONT, HEAD_FONT, MONO_FONT = "Cambria", "Arial", "Courier New"
C_INK, C_INDIGO, C_MUTED, C_BODY = RGBColor(0x14, 0x19, 0x36), RGBColor(0x3B, 0x46, 0xC4), RGBColor(0x5B, 0x5F, 0x77), RGBColor(0x1E, 0x22, 0x33)
C_AMBER = RGBColor(0xB8, 0x74, 0x00)
TEXT_WIDTH_CM = 16.0
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def _set_font(target, name: str) -> None:
    """Set a font on a style or run for every script, removing theme-font overrides."""
    rpr = target.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        if attr.endswith("Theme"):
            del fonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _runs(paragraph, text: str, size: float | None = None, color: RGBColor | None = None, font: str | None = None, bold: bool | None = None, italic: bool | None = None):
    for piece in INLINE.split(text):
        if not piece:
            continue
        run_bold, run_italic, mono = bold, italic, False
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, run_bold = piece[2:-2], True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            piece, mono = piece[1:-1], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, run_italic = piece[1:-1], True
        run = paragraph.add_run(piece)
        if run_bold is not None:
            run.bold = run_bold
        if run_italic is not None:
            run.italic = run_italic
        if mono:
            _set_font(run, MONO_FONT)
            run.font.size = Pt((size or 10.5) - 0.5)
        else:
            if font:
                _set_font(run, font)
            if size:
                run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _cell_borders(cell, **edges: tuple[int, str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            size, color = edges[edge]
            el.set(qn("w:val"), "single"), el.set(qn("w:sz"), str(size)), el.set(qn("w:space"), "0"), el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)


def _table_setup(tbl, widths_cm: list[float], margins: tuple[int, int] = (45, 90)) -> None:
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl_pr = tbl._tbl.tblPr
    look = tbl_pr.find(qn("w:tblLook"))  # autofit = False has already written a fixed w:tblLayout
    place = (lambda el: look.addprevious(el)) if look is not None else tbl_pr.append
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", margins[0]), ("left", margins[1]), ("bottom", margins[0]), ("right", margins[1])):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value)), el.set(qn("w:type"), "dxa")
        mar.append(el)
    place(mar)
    for row in tbl.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
    grid = tbl._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_cm):
        col.set(qn("w:w"), str(int(width / 2.54 * 1440)))


def _row_flags(row, header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruction
            run._r.append(el)
        elif kind == "text":
            run.text = placeholder
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def _fld_char(paragraph, kind: str, instruction: str | None = None) -> None:
    run = paragraph.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)
    if instruction:
        run = paragraph.add_run()
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = instruction
        run._r.append(el)


# Children of w:pPr / w:rPr that must come after the element being inserted (the schema fixes their order).
PPR_AFTER_BORDER = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
                    "w:rPr", "w:sectPr", "w:pPrChange")
RPR_AFTER_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                     "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _bottom_rule(paragraph, color: str = "F2A51A", size: int = 8, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"), bottom.set(qn("w:sz"), str(size)), bottom.set(qn("w:space"), str(space)), bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.insert_element_before(borders, *PPR_AFTER_BORDER)


def _hyperlink(paragraph, url: str, size: float = 10.5) -> None:
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = paragraph.add_run(url)
    _set_font(run, HEAD_FONT)
    run.font.size, run.font.color.rgb, run.font.underline = Pt(size), C_INDIGO, True
    link.append(run._r)
    paragraph._p.append(link)


def _spacing(paragraph, before: float = 0, after: float = 6, line: float | None = None, keep: bool = False):
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    if line:
        fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True
    return paragraph


def _styles(doc) -> None:
    normal = doc.styles["Normal"]
    _set_font(normal, BODY_FONT)
    normal.font.size, normal.font.color.rgb = Pt(10.5), C_BODY
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(7), 1.22
    normal.paragraph_format.widow_control = True
    for name, size, color, before, after in (("Heading 1", 17, C_INK, 22, 9), ("Heading 2", 12, C_INDIGO, 13, 4), ("Title", 34, C_INK, 0, 6)):
        style = doc.styles[name]
        _set_font(style, HEAD_FONT)
        style.font.size, style.font.bold, style.font.italic, style.font.color.rgb = Pt(size), True, False, color
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.keep_with_next, style.paragraph_format.line_spacing = True, 1.1
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    caption = doc.styles["Caption"]
    _set_font(caption, HEAD_FONT)
    caption.font.size, caption.font.bold, caption.font.italic, caption.font.color.rgb = Pt(8.5), False, False, C_MUTED
    caption.paragraph_format.line_spacing = 1.15
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        _set_font(style, BODY_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
    for level, indent, size, bold in ((1, 0.0, 9.5, True), (2, 0.6, 9, False)):
        try:
            style = doc.styles[f"TOC {level}"]
        except KeyError:
            style = doc.styles.add_style(f"TOC {level}", 1)
        style.base_style = normal
        _set_font(style, HEAD_FONT)
        style.font.size, style.font.bold, style.font.color.rgb = Pt(size), bold, C_INK if bold else C_BODY
        fmt = style.paragraph_format
        fmt.left_indent, fmt.space_before, fmt.space_after, fmt.line_spacing = Cm(indent), Pt(5 if bold else 0), Pt(1), 1.05
        fmt.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def _title_page(doc) -> None:
    label = _spacing(doc.add_paragraph(), before=96, after=10)
    run = label.add_run("APPLIED AI & MACHINE LEARNING CAPSTONE")
    _set_font(run, HEAD_FONT)
    run.font.size, run.font.bold, run.font.color.rgb = Pt(9.5), True, C_AMBER
    run._r.get_or_add_rPr().insert_element_before(_spacing_el(40), *RPR_AFTER_SPACING)
    doc.add_paragraph("AI Portfolio Advisor", style="Title")
    sub = _spacing(doc.add_paragraph(), after=14, line=1.25)
    _runs(sub, SUBTITLE, size=14, color=C_MUTED, italic=True)
    _bottom_rule(sub, size=12, space=14)
    kind = _spacing(doc.add_paragraph(), before=14, after=2)
    _runs(kind, "Capstone project report", size=13, color=C_INK, font=HEAD_FONT, bold=True)
    _runs(_spacing(doc.add_paragraph(), after=34), COURSE, size=10.5, color=C_MUTED, font=HEAD_FONT)

    rows = [("Team", GROUP)] + [("", f"{name}   ·   {roll}") for name, roll in TEAM] + [("Live application", APP_URL), ("Landing page", LANDING_URL), ("Source code", REPO_URL),
                                                                                         ("Date", pd.Timestamp(M["generated"]).strftime("%B %Y"))]
    tbl = doc.add_table(rows=len(rows), cols=2)
    _table_setup(tbl, [3.6, 12.4], margins=(50, 0))
    for row, (key, value) in zip(tbl.rows, rows):
        left, right = row.cells[0].paragraphs[0], row.cells[1].paragraphs[0]
        _spacing(left, after=0), _spacing(right, after=0)
        _runs(left, key.upper(), size=8, color=C_MUTED, font=HEAD_FONT, bold=True)
        if value.startswith("http"):
            _hyperlink(right, value, 10)
        else:
            _runs(right, value, size=10.5 if key != "Team" else 11, color=C_INK, font=HEAD_FONT, bold=(key == "Team"))
    note = _spacing(doc.add_paragraph(), before=70, after=0, line=1.2)
    _runs(note, DISCLAIMER + f" Every figure and number in this report is generated by `scripts/build_report.py` from `reports/metrics.json` (model run of {M['generated']}).",
          size=8.5, color=C_MUTED, font=HEAD_FONT)


def _spacing_el(value: int):
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(value))
    return el


def _toc(doc, pages: dict[str, int] | None) -> None:
    heading = doc.add_paragraph("Contents", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    entries = [(1 if b["kind"] == "h1" else 2, b["text"]) for b in BLOCKS if b["kind"] in ("h1", "h2")]
    for i, (level, text) in enumerate(entries):
        para = doc.add_paragraph(style=f"TOC {level}")
        if i == 0:
            _fld_char(para, "begin", ' TOC \\o "1-2" \\h \\z \\u ')
            _fld_char(para, "separate")
        para.add_run(f"{text}\t{(pages or {}).get(text, 0) or ''}" if pages else f"{text}\t0")
        if i == len(entries) - 1:
            _fld_char(para, "end")


def _add_table(doc, b: dict) -> None:
    caption = doc.add_paragraph(style="Caption")
    _spacing(caption, before=8, after=4, keep=True)
    _runs(caption, f"**Table {b['number']}.** {b['caption']}", size=8.5, color=C_MUTED, font=HEAD_FONT)
    size = b["size"]
    tbl = doc.add_table(rows=1 + len(b["rows"]), cols=len(b["header"]))
    _table_setup(tbl, b["widths"])
    for r, row in enumerate(tbl.rows):
        values = b["header"] if r == 0 else b["rows"][r - 1]
        _row_flags(row, header=(r == 0))
        for c, cell in enumerate(row.cells):
            para = cell.paragraphs[0]
            _spacing(para, after=0, line=1.12)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if b["align"][c] == "r" else WD_ALIGN_PARAGRAPH.LEFT
            if r == 0:
                _runs(para, values[c], size=size, color=RGBColor(0xFF, 0xFF, 0xFF), font=HEAD_FONT, bold=True)
                _cell_borders(cell)
                _shade(cell, "141936")
                para.paragraph_format.keep_with_next = True
            else:
                _runs(para, values[c] if c < len(values) else "", size=size, color=C_BODY, font=HEAD_FONT)
                _cell_borders(cell, bottom=(4, "DADCE6"))
                if r % 2 == 0:
                    _shade(cell, "F6F7FB")
                if r <= 1 and len(b["rows"]) > 2:
                    para.paragraph_format.keep_with_next = True
    _spacing(doc.add_paragraph(), after=4).paragraph_format.line_spacing = 0.6


def _add_box(doc, fill: str, bar: str | None):
    tbl = doc.add_table(rows=1, cols=1)
    _table_setup(tbl, [TEXT_WIDTH_CM], margins=(110, 170))
    cell = tbl.rows[0].cells[0]
    _cell_borders(cell, **({"left": (24, bar)} if bar else {}))
    _shade(cell, fill)
    _row_flags(tbl.rows[0])
    return cell


def render_docx(figs: dict[str, Path], pages: dict[str, int] | None, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.3)
    section.different_first_page_header_footer = True
    _styles(doc)
    for zoom in doc.settings.element.findall(qn("w:zoom")):
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = "AI Portfolio Advisor: Capstone Report"
    doc.core_properties.author = f"{GROUP}: " + ", ".join(name for name, _ in TEAM)
    doc.core_properties.subject = COURSE
    doc.core_properties.keywords = "quantile regression, LightGBM, SHAP, walk-forward cross-validation, portfolio"

    doc.styles["Footer"].paragraph_format.tab_stops.clear_all()
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    _runs(footer, f"AI Portfolio Advisor  ·  Capstone report  ·  {GROUP}\t", size=8, color=C_MUTED, font=HEAD_FONT)
    start = len(footer.runs)
    _field(footer, " PAGE ", "1")
    for run in footer.runs[start:]:
        _set_font(run, HEAD_FONT)
        run.font.size, run.font.color.rgb = Pt(8), C_MUTED

    for b in BLOCKS:
        kind = b["kind"]
        if kind == "title":
            _title_page(doc)
        elif kind == "toc":
            _toc(doc, pages)
        elif kind == "h1":
            para = doc.add_paragraph(b["text"], style="Heading 1")
            para.paragraph_format.page_break_before = bool(b["page_break"])
            _bottom_rule(para, size=6, space=5)
        elif kind == "h2":
            doc.add_paragraph(b["text"], style="Heading 2")
        elif kind == "p":
            _runs(doc.add_paragraph(), b["text"])
        elif kind == "ref":
            para = _runs(doc.add_paragraph(), b["text"], size=9.5)
            para.paragraph_format.left_indent, para.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
            _spacing(para, after=4, line=1.15)
        elif kind in ("bullets", "numbered"):
            for item in b["items"]:
                _runs(doc.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number"), item)
        elif kind == "table":
            _add_table(doc, b)
        elif kind == "figure":
            para = _spacing(doc.add_paragraph(), before=6, after=3, keep=True)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(figs[b["key"]]), width=Cm(b["width_cm"]))
            caption = doc.add_paragraph(style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(caption, after=10)
            _runs(caption, f"**Figure {b['number']}.** {b['caption']}", size=8.5, color=C_MUTED, font=HEAD_FONT)
        elif kind == "callout":
            cell = _add_box(doc, "FFF7E6", "F2A51A")
            head = _spacing(cell.paragraphs[0], after=3, keep=True)
            _runs(head, b["title"], size=9.5, color=C_INK, font=HEAD_FONT, bold=True)
            _runs(_spacing(cell.add_paragraph(), after=0, line=1.2), b["text"], size=9.5, color=C_BODY, font=HEAD_FONT)
            _spacing(doc.add_paragraph(), after=4).paragraph_format.line_spacing = 0.6
        elif kind == "eq":
            para = _spacing(doc.add_paragraph(), before=2, after=8)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(para, b["text"], size=10.5, color=C_INK, italic=True)
        elif kind == "code":
            cell = _add_box(doc, "F4F5F9", None)
            for i, line in enumerate(b["lines"]):
                para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                _spacing(para, after=0, line=1.15)
                run = para.add_run(line)
                _set_font(run, MONO_FONT)
                run.font.size, run.font.color.rgb = Pt(8.5), C_INK
            _spacing(doc.add_paragraph(), after=4).paragraph_format.line_spacing = 0.6
        elif kind == "qa":
            question = _spacing(doc.add_paragraph(), before=5, after=2, keep=True)
            _runs(question, f"Q. {b['q']}", size=9.5, color=C_INK, font=HEAD_FONT, bold=True)
            _runs(_spacing(doc.add_paragraph(), after=4, line=1.18), b["a"], size=9.5)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


# -------------------------------------------------------------------- PDF
def _soffice() -> str | None:
    return shutil.which("soffice") or next((p for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",) if Path(p).exists()), None)


def convert_to_pdf(docx: Path, out_dir: Path) -> Path | None:
    exe = _soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as profile:
        result = subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx)],
                                capture_output=True, text=True, timeout=300, check=False)
    pdf = out_dir / (docx.stem + ".pdf")
    if not pdf.exists():
        print("PDF conversion failed:", result.stderr.strip()[:300])
        return None
    return pdf


def heading_pages(pdf: Path) -> dict[str, int] | None:
    """Page number of every heading, read back from the rendered PDF."""
    if not shutil.which("pdftotext"):
        return None
    text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, check=False).stdout
    page_lines = [[" ".join(line.split()) for line in page.splitlines()] for page in text.split("\f")]
    pages: dict[str, int] = {}
    cursor = 0
    for b in BLOCKS:
        if b["kind"] not in ("h1", "h2"):
            continue
        target = " ".join(b["text"].split())
        for number in range(cursor, len(page_lines)):
            lines = page_lines[number]
            if target in lines or any(line and target.startswith(line) and len(line) > 25 for line in lines):
                pages[b["text"]] = number + 1
                cursor = number
                break
    return pages


def document_stats(path: Path) -> dict[str, int]:
    doc = Document(path)
    paragraphs = list(doc.paragraphs)
    words = sum(len(p.text.split()) for p in paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                words += len(cell.text.split())
    return {
        "headings": sum(1 for p in paragraphs if p.style.name.startswith("Heading")),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "words": words,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--no-pdf", action="store_true", help="skip the PDF and the page-numbered table of contents")
    args = parser.parse_args()

    for warning in check_narrative():
        print(f"WARNING: narrative out of date: {warning}")
    figs = build_figures()
    build_content()
    resolve_references()
    MD_PATH.write_text(render_markdown(figs))

    pages = None
    if not args.no_pdf and _soffice():
        with tempfile.TemporaryDirectory() as tmp:
            draft = Path(tmp) / DOCX_PATH.name
            render_docx(figs, None, draft)  # first pass: lay the document out to learn the page numbers
            draft_pdf = convert_to_pdf(draft, Path(tmp))
            pages = heading_pages(draft_pdf) if draft_pdf else None
    render_docx(figs, pages, DOCX_PATH)
    pdf = None if args.no_pdf else convert_to_pdf(DOCX_PATH, DOCS)

    stats = document_stats(DOCX_PATH)
    print(f"Wrote {DOCX_PATH.relative_to(ROOT)}: {stats['headings']} headings, {stats['tables']} tables, {stats['images']} images, about {stats['words']:,} words")
    print(f"Wrote {MD_PATH.relative_to(ROOT)} and {len(figs)} figures in {FIGS.relative_to(ROOT)}/")
    if pdf:
        print(f"Wrote {pdf.relative_to(ROOT)}" + ("" if pages else " (table of contents has no page numbers: pdftotext not found)"))
    elif not args.no_pdf:
        print("No PDF: LibreOffice (soffice) not found. Open the .docx in Word and update the table of contents (right-click, Update Field).")


if __name__ == "__main__":
    main()
