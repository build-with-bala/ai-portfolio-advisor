"""Build the written report (Markdown + Word) and the slide deck from reports/metrics.json.

    python scripts/build_submission.py
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
R = json.loads((ROOT / "reports" / "metrics.json").read_text())
D, SEL, HOLD, BT = R["data"], R["selected"], R["holdout"], R["backtest"]
LG, NV, RG = HOLD["lightgbm"], HOLD["naive"], HOLD["ridge"]
CAND = {c["model"]: c for c in R["candidates"]}
TUNED, NOTEBOOK, NAIVE = CAND["LightGBM quantile, tuned"], CAND["LightGBM quantile, notebook settings"], CAND["Naive (historical quantiles)"]
APP_URL = "https://fn.iimbg.com"
REPO_URL = "https://github.com/build-with-bala/ai-portfolio-advisor"

# Adviser-productivity case. Every input here is an assumption, not a measurement.
ROI = dict(clients=200, reviews_per_year=4, minutes_before=45, minutes_after=10, adviser_rate_inr=1500, annual_cost_inr=120000)
hours_saved = ROI["clients"] * ROI["reviews_per_year"] * (ROI["minutes_before"] - ROI["minutes_after"]) / 60
value = hours_saved * ROI["adviser_rate_inr"]
roi = (value - ROI["annual_cost_inr"]) / ROI["annual_cost_inr"]


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


params = ", ".join(f"{k}={v:g}" for k, v in SEL["params"].items())
top_shap = ", ".join(s["feature"] for s in R["shap"][:4])
h, cv = BT["holdout"], BT["cross_validation_folds"]

SECTIONS: list[tuple[str, list[str]]] = [
    ("1. Problem and user", [
        "Retail investors and small wealth advisers in India pick stocks from tips and single-number targets, with no view of the range of outcomes or of how a pick fits their risk tolerance.",
        "Target user: an independent wealth adviser or self-directed investor who reviews a portfolio monthly.",
        "AI Portfolio Advisor takes a risk profile and capital, and returns a ranked book of NSE stocks marked STRONG BUY / BUY / WATCH / PASS, each with a pessimistic, median and optimistic 21-day return, the evidence behind it, and a position size.",
        f"Working prototype: {APP_URL}  |  Code: {REPO_URL}",
    ]),
    ("2. Business value and ROI", [
        f"Time saved (assumptions, editable in scripts/build_submission.py): {ROI['clients']} clients x {ROI['reviews_per_year']} reviews a year, screening time cut from {ROI['minutes_before']} to {ROI['minutes_after']} minutes, adviser time valued at Rs {ROI['adviser_rate_inr']:,}/hour.",
        f"That is {hours_saved:,.0f} adviser hours a year, worth Rs {value:,.0f}, against an assumed running cost of Rs {ROI['annual_cost_inr']:,} a year: a return of {pct(roi, 0)} on cost.",
        f"Risk communication: {pct(LG['coverage'])} of real 21-day outcomes on unseen data fell inside the model's P10-P90 range, against a target of {pct(HOLD['target_coverage'], 0)}. A client can be shown an honest range.",
        f"What we do not claim: extra return. In the hold-out backtest the top-5 strategy returned {pct(h['strategy_annual_return_median'])} a year against {pct(h['benchmark_annual_return_median'])} for holding all stocks equally (median over {h['runs']} start days, after costs). The product is sold on speed, discipline and honest risk ranges, not on beating the market.",
        "Market strategy: free research tool for self-directed investors; paid seat for advisers who need client-ready evidence sheets; brokers as a channel through the existing Upstox / Zerodha live-quote adapters.",
    ]),
    ("3. System architecture", [
        "src/data.py: download prices (Yahoo Finance), build 19 candidate features per stock per day, select features, define time-series splits.",
        "src/model.py: tune and cross-validate, train three LightGBM quantile models, evaluate on a hold-out period, compute SHAP, run a backtest, export data/artifacts.pkl and reports/metrics.json.",
        "src/portfolio_advisor/core.py: decision engine. Combines the technical score, fundamental score, model forecast and risk fit by profile weights, sizes positions with Hierarchical Risk Parity, a position cap and whole shares.",
        "app.py: Streamlit interface (profile intake, ranked book, stock research with SHAP, portfolio lab, validation, model report, live NSE quotes). api_app.py: optional FastAPI service.",
        "Flow: prices -> features -> selection -> purged walk-forward CV -> quantile models -> artifact -> app.",
    ]),
    ("4. Data and preprocessing", [
        f"{D['tickers']} liquid NSE large caps, five per sector across {D['sectors']} sectors, daily adjusted prices {D['first_date']} to {D['last_date']}: {D['panel_rows']:,} stock-days.",
        f"Label: return over the next {R['config']['HORIZON']} trading days.",
        "Features are returns, ratios and oscillators (RSI, MACD histogram scaled by price, Bollinger bandwidth, ATR, stochastic %K, OBV change, Williams %R, momentum, volatility, distance from 52-week high, volume surge, market-relative return, market return and volatility).",
        R["preprocessing"]["stationarity"],
        R["preprocessing"]["scaling"],
        R["preprocessing"]["missing_values"],
    ]),
    ("5. Feature selection", [
        f"Method: {R['feature_selection']['method']}, run on development data only.",
        f"Kept {len(R['feature_selection']['selected'])} of {len(R['feature_selection']['candidates'])}: {', '.join(R['feature_selection']['selected'])}.",
        "Why: many indicators measure the same thing (RSI, stochastic %K and Williams %R all track short-term position in the range). The redundancy penalty keeps one of each kind.",
        "Company fundamentals were removed from the model. Only today's ratios are available, so using them on past dates would leak future information. They still feed the app's separate fundamental score.",
    ]),
    ("6. Tuning and cross-validation", [
        f"Scheme: {R['validation']['scheme']}. Each fold trains on the past and validates on the next block, with {R['config']['HORIZON'] + R['config']['EMBARGO']} trading days removed in between because each label looks {R['config']['HORIZON']} days ahead.",
        f"Search: random search over {R['validation']['lightgbm_trials']} settings x {R['config']['N_SPLITS']} folds x 3 quantiles = {R['validation']['lightgbm_fits']} model fits. Tuning metric: mean pinball loss (the loss a quantile forecast is meant to minimise).",
        f"Chosen settings: {params}.",
        f"CV pinball loss: tuned {TUNED['pinball_mean']:.5f} +/- {TUNED['pinball_std']:.5f}; notebook settings {NOTEBOOK['pinball_mean']:.5f}; naive baseline {NAIVE['pinball_mean']:.5f}.",
        f"Tuning improved on the original notebook settings by {pct(SEL['pinball_improvement_vs_notebook'])}. The search picked very small trees ({SEL['params']['num_leaves']} leaves, {SEL['params']['n_estimators']} trees): with a weak signal, complex trees fit noise.",
        f"Against the naive no-feature baseline the tuned model is {pct(abs(SEL['pinball_improvement_vs_naive']))} {'better' if SEL['pinball_improvement_vs_naive'] > 0 else 'worse'} in cross-validation, which is inside fold-to-fold variation. We report this as it is.",
    ]),
    ("7. Hold-out results", [
        f"Hold-out period: {D['holdout_start']} to {D['holdout_end_scored']}, {D['holdout_rows_scored']:,} stock-days never used for selection or tuning.",
        f"Pinball loss: LightGBM {LG['pinball']:.5f} vs naive {NV['pinball']:.5f}. MAE: LightGBM {LG['mae']:.4f}, Ridge {RG['mae']:.4f}, naive {NV['mae']:.4f}.",
        f"Rank information coefficient (daily rank correlation of forecast and outcome): LightGBM {LG['ic_mean']:+.3f}, positive on {pct(LG['ic_positive_share'], 0)} of {LG['ic_days']} days; Ridge {RG['ic_mean']:+.3f}.",
        f"Direction right {pct(LG['directional_accuracy'])} of the time. P10-P90 coverage {pct(LG['coverage'])} against a {pct(HOLD['target_coverage'], 0)} target.",
        f"Backtest, hold-out: strategy {pct(h['strategy_annual_return_median'])}/yr vs benchmark {pct(h['benchmark_annual_return_median'])}/yr; excess {pct(h['annual_excess_median'])} (range {pct(h['annual_excess_min'])} to {pct(h['annual_excess_max'])}); beat the benchmark from {pct(h['share_of_runs_beating_benchmark'], 0)} of start days.",
        f"Backtest, cross-validation folds: excess {pct(cv['annual_excess_median'])}/yr (range {pct(cv['annual_excess_min'])} to {pct(cv['annual_excess_max'])}).",
    ]),
    ("8. Explainability and error analysis", [
        f"SHAP (TreeExplainer on the median model). Largest drivers: {top_shap}.",
        "Every stock page in the app shows that stock's own SHAP drivers for the current forecast.",
        *[f"{row['group'].capitalize()} markets: MAE {row['mae']:.4f}, coverage {pct(row['coverage'])}, rank IC {row['ic_mean']:+.3f}." for row in R["error_analysis"]["volatility_regime"]],
        *[f"When the {row['group']}: MAE {row['mae']:.4f}, direction right {pct(row['directional_accuracy'])}." for row in R["error_analysis"]["market_direction"]],
    ]),
    ("9. Limits and next steps", [*R["limits"], "Next: point-in-time fundamentals, a market-relative target, a larger universe including delisted names, and a controlled paper-trading period."]),
    ("10. Repository and reproducibility", [
        "Layout: src/data.py, src/model.py, app.py, requirements.txt, tests/ (17 tests), reports/, docs/.",
        "Reproduce: pip install -r requirements.txt; python -m src.model (trains from the committed data snapshot in about a minute); streamlit run app.py.",
        "Vibe Coding Log: docs/VIBE_CODING_LOG.md.",
    ]),
]

TITLE = "AI Portfolio Advisor"
SUB = "Applied AI & Machine Learning Capstone | AI and ML for Digital Business Managers"

md = [f"# {TITLE}", "", f"*{SUB}*", ""]
for head, lines in SECTIONS:
    md += [f"## {head}", ""] + [f"- {line}" for line in lines] + [""]
(ROOT / "docs" / "REPORT.md").write_text("\n".join(md))

doc = Document()
doc.add_heading(TITLE, 0)
doc.add_paragraph(SUB)
for head, lines in SECTIONS:
    doc.add_heading(head, 1)
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")
doc.save(ROOT / "docs" / "Capstone_Report.docx")

NAVY, INK, ACCENT = RGBColor(0x14, 0x1B, 0x3C), RGBColor(0x22, 0x22, 0x22), RGBColor(0xE0, 0x9F, 0x1F)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def slide(title: str, lines: list[str], tag: str = "") -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.25))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(11), Inches(0.8)).text_frame
    box.text = title
    box.paragraphs[0].font.size, box.paragraphs[0].font.bold, box.paragraphs[0].font.color.rgb = Pt(30), True, RGBColor(255, 255, 255)
    if tag:
        t = s.shapes.add_textbox(Inches(10.3), Inches(0.42), Inches(2.6), Inches(0.5)).text_frame
        t.text = tag
        t.paragraphs[0].font.size, t.paragraphs[0].font.color.rgb = Pt(14), ACCENT
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.5)).text_frame
    body.word_wrap = True
    size = Pt(20) if len(lines) <= 5 else Pt(17)
    for i, line in enumerate(lines):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = "•  " + line
        p.font.size, p.font.color.rgb, p.space_after = size, INK, Pt(12)


cover = prs.slides.add_slide(prs.slide_layouts[6])
bg = cover.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tf = cover.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.5), Inches(3)).text_frame
tf.text = TITLE
tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold, tf.paragraphs[0].font.color.rgb = Pt(54), True, RGBColor(255, 255, 255)
for text, size, color in [("A forecast is a range: profile-aware stock recommendations with visible uncertainty", 24, ACCENT), (SUB, 16, RGBColor(220, 220, 230)), (f"{APP_URL}   |   {REPO_URL}", 16, RGBColor(220, 220, 230))]:
    p = tf.add_paragraph(); p.text = text; p.font.size = Pt(size); p.font.color.rgb = color; p.space_before = Pt(14)

TAGS = ["Executive pitch", "Executive pitch", "Technical deep dive", "Technical deep dive", "Technical deep dive", "Technical deep dive", "Technical deep dive", "Technical deep dive", "Technical deep dive", "Repository"]
for (head, lines), tag in zip(SECTIONS, TAGS):
    slide(head.split(". ", 1)[1], lines, tag)
slide("Live demonstration", [
    f"Open {APP_URL}.",
    "Set a Conservative profile with Rs 5,00,000: read the ranked book and the cash remainder.",
    "Switch to Aggressive: the same evidence, different weights and caps, different book.",
    "Open one stock: P10 / P50 / P90 range and its SHAP drivers.",
    "Model report tab: cross-validation table, hold-out results, backtest, error analysis.",
], "Demo")
prs.save(ROOT / "docs" / "Capstone_Deck.pptx")
print("Built docs/REPORT.md, docs/Capstone_Report.docx, docs/Capstone_Deck.pptx")
print(f"ROI case: {hours_saved:.0f} h, Rs {value:,.0f}, ROI {roi:.0%}")
